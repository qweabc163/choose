#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
A股趋势策略：指定日期股票分类 + 未来走势回测 + 动态卖出验证

特点：
1. 历史行情只从本地 data/*.csv 读取；
2. 沪深300基准只从 data/_benchmark_000300.csv 读取；
3. 默认不联网；
4. --universe-source ak 只联网获取股票代码和名称，不下载历史行情；
5. 自动剔除 ST / *ST / 退市风险；
6. 仅保留沪深主板：60xxxx / 00xxxx；
7. 支持指定日期股票分类；
8. 支持固定持有 N 日验证；
9. 支持动态卖出验证；
10. 输出 Excel。


 python pool03.py --date 2026-04-30 --classify-only
"""

import os
import re
import argparse
import warnings
from typing import Dict, List, Optional

import numpy as np
import pandas as pd
from tqdm import tqdm
from numpy.lib.stride_tricks import sliding_window_view

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


# ==============================
# 0. 默认参数
# ==============================

CACHE_DIR = "data"
OUTPUT_DIR = "output"

BENCHMARK_CACHE_FILE = os.path.join(CACHE_DIR, "_benchmark_000300.csv")
STOCK_NAME_FILE = os.path.join(CACHE_DIR, "_stock_names.csv")

MIN_SCORE = 75
POOL_STATUSES = ["强趋势池", "趋势观察池"]

MIN_BARS = 130
MIN_AVG_AMOUNT_20 = 80_000_000

BUY_PRICE_FIELD = "open"

FORWARD_ANALYSIS_OUTPUT_PREFIX = "pool_forward"
CLASSIFICATION_OUTPUT_PREFIX = "pool_classification"

DYNAMIC_EXIT_ENABLED = True
MIN_HOLD_DAYS = 3
EXIT_SELL_TIMING = "next_open"
STOP_LOSS_PCT = -0.08
TRAILING_START_PCT = 0.10
TRAILING_DRAWDOWN_PCT = 0.08


# ==============================
# 1. 工具函数
# ==============================

def normalize_code(code) -> str:
    s = str(code).strip()
    m = re.search(r"(\d{6})", s)
    if m:
        return m.group(1)
    return s.zfill(6)


def is_mainboard_code(code: str) -> bool:
    code = normalize_code(code)
    return len(code) == 6 and code.isdigit() and code.startswith(("60", "00"))


def is_bad_name(name: str) -> bool:
    if pd.isna(name):
        return False
    s = str(name).strip()
    if not s:
        return False
    return bool(re.search(r"ST|\*ST|退", s, flags=re.IGNORECASE))


def max_drawdown(series: pd.Series) -> float:
    s = series.dropna()
    if len(s) == 0:
        return np.nan
    dd = s / s.cummax() - 1
    return float(dd.min())


def rolling_max_drawdown_np(close: pd.Series, window: int = 20) -> np.ndarray:
    arr = close.to_numpy(dtype="float64", copy=False)
    n = len(arr)
    out = np.full(n, np.nan)

    if n < window:
        return out

    windows = sliding_window_view(arr, window_shape=window)
    valid = ~np.isnan(windows).any(axis=1)
    vals = np.full(windows.shape[0], np.nan)

    if valid.any():
        w = windows[valid]
        cum_max = np.maximum.accumulate(w, axis=1)
        dd = w / cum_max - 1
        vals[valid] = dd.min(axis=1)

    out[window - 1:] = vals
    return out


def get_row_price(row: pd.Series, field: str, fallback: str = "close") -> float:
    if field in row.index and not pd.isna(row[field]):
        return float(row[field])
    if fallback in row.index and not pd.isna(row[fallback]):
        return float(row[fallback])
    return np.nan


def normalize_date_col(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["date"] = pd.to_datetime(df["date"], errors="coerce").dt.normalize()
    df = df.dropna(subset=["date"])
    df = df.sort_values("date").drop_duplicates("date", keep="last").reset_index(drop=True)
    return df


def standardize_local_df(df: pd.DataFrame, code: str = "") -> Optional[pd.DataFrame]:
    if df is None or df.empty:
        return None

    df = df.copy()

    rename_map = {
        "日期": "date",
        "时间": "date",
        "开盘": "open",
        "收盘": "close",
        "最高": "high",
        "最低": "low",
        "成交量": "volume",
        "成交额": "amount",
        "成交金额": "amount",
        "涨跌幅": "pct_chg",
        "涨跌额": "change",
        "换手率": "turnover",
        "振幅": "amplitude",
        "date": "date",
        "open": "open",
        "close": "close",
        "high": "high",
        "low": "low",
        "volume": "volume",
        "amount": "amount",
        "pct_chg": "pct_chg",
        "turnover": "turnover",
        "change": "change",
        "amplitude": "amplitude",
    }

    df = df.rename(columns=rename_map)

    if "date" not in df.columns:
        print(f"{code} 缺少 date 字段，跳过。")
        return None

    required_cols = ["date", "open", "high", "low", "close"]
    missing = [c for c in required_cols if c not in df.columns]

    if missing:
        print(f"{code} 缺少必要字段 {missing}，跳过。")
        return None

    df = normalize_date_col(df)

    numeric_cols = [
        "open", "high", "low", "close", "volume", "amount",
        "pct_chg", "turnover", "change", "amplitude",
    ]

    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df = df.dropna(subset=["open", "high", "low", "close"])

    if df.empty:
        return None

    if "volume" not in df.columns:
        df["volume"] = np.nan

    if "amount" not in df.columns or df["amount"].isna().all():
        if "volume" in df.columns and not df["volume"].isna().all():
            df["amount"] = df["volume"] * df["close"]
        else:
            df["amount"] = np.nan

    if "pct_chg" not in df.columns or df["pct_chg"].isna().all():
        df["pct_chg"] = df["close"].pct_change() * 100

    if "turnover" not in df.columns:
        df["turnover"] = np.nan

    if "change" not in df.columns:
        df["change"] = df["close"].diff()

    return df.reset_index(drop=True)


# ==============================
# 2. 股票列表和名称
# ==============================

def scan_local_stock_codes(data_dir: str = "data") -> List[str]:
    if not os.path.exists(data_dir):
        raise FileNotFoundError(f"数据目录不存在：{data_dir}")

    codes = []

    for fname in os.listdir(data_dir):
        if not fname.lower().endswith(".csv"):
            continue

        stem = fname[:-4]

        if stem.startswith("_benchmark") or stem.startswith("_stock_names"):
            continue

        code = normalize_code(stem)

        if is_mainboard_code(code):
            codes.append(code)

    return sorted(set(codes))


def load_local_stock_names(data_dir: str = "data") -> Dict[str, str]:
    path = os.path.join(data_dir, "_stock_names.csv")

    if not os.path.exists(path):
        return {}

    try:
        df = pd.read_csv(path, dtype=str)
    except Exception as e:
        print(f"读取本地股票名称文件失败：{path}，错误：{e}")
        return {}

    if df is None or df.empty:
        return {}

    code_col = None
    name_col = None

    for c in ["code", "代码", "股票代码", "symbol", "证券代码"]:
        if c in df.columns:
            code_col = c
            break

    for c in ["name", "名称", "股票简称", "stock_name", "证券简称"]:
        if c in df.columns:
            name_col = c
            break

    if code_col is None or name_col is None:
        print(f"本地股票名称文件字段异常：{path}")
        return {}

    tmp = df[[code_col, name_col]].copy()
    tmp.columns = ["code", "name"]
    tmp["code"] = tmp["code"].apply(normalize_code)
    tmp["name"] = tmp["name"].astype(str).str.strip()
    tmp = tmp[tmp["code"].apply(is_mainboard_code)]
    tmp = tmp[~tmp["name"].apply(is_bad_name)]
    tmp = tmp.drop_duplicates(subset=["code"], keep="last")

    return dict(zip(tmp["code"], tmp["name"]))


def save_stock_names(data_dir: str, df: pd.DataFrame) -> None:
    if df is None or df.empty:
        return

    os.makedirs(data_dir, exist_ok=True)
    path = os.path.join(data_dir, "_stock_names.csv")

    out = df[["code", "name"]].copy()
    out["code"] = out["code"].astype(str).apply(normalize_code)
    out["name"] = out["name"].astype(str).str.strip()
    out = out[out["code"].apply(is_mainboard_code)]
    out = out[~out["name"].apply(is_bad_name)]
    out = out.drop_duplicates(subset=["code"], keep="last")
    out = out.sort_values("code").reset_index(drop=True)

    try:
        out.to_csv(path, index=False, encoding="utf-8-sig")
        print(f"股票名称已保存到：{path}")
    except Exception as e:
        print(f"保存股票名称失败：{e}")


def get_a_stock_universe_from_ak() -> pd.DataFrame:
    print("正在通过 AkShare 获取 A 股股票列表和名称...")

    try:
        import akshare as ak
    except Exception as e:
        raise RuntimeError(
            "未安装 AkShare。如需使用 --universe-source ak，请先执行：pip install akshare"
        ) from e

    stock_info = None
    last_err = None

    try:
        stock_info = ak.stock_info_a_code_name()
    except Exception as e:
        last_err = e
        stock_info = None

    if stock_info is None or stock_info.empty:
        try:
            stock_info = ak.stock_zh_a_spot()
        except Exception as e:
            last_err = e
            stock_info = None

    if stock_info is None or stock_info.empty:
        raise RuntimeError(f"无法通过 AkShare 获取股票列表，请检查网络。最后错误：{last_err}")

    if "code" in stock_info.columns and "name" in stock_info.columns:
        df = stock_info[["code", "name"]].copy()
    elif "代码" in stock_info.columns and "名称" in stock_info.columns:
        df = stock_info[["代码", "名称"]].copy()
        df.columns = ["code", "name"]
    elif "symbol" in stock_info.columns and "name" in stock_info.columns:
        df = stock_info[["symbol", "name"]].copy()
        df.columns = ["code", "name"]
    elif "代码" in stock_info.columns and "简称" in stock_info.columns:
        df = stock_info[["代码", "简称"]].copy()
        df.columns = ["code", "name"]
    else:
        raise ValueError(f"股票列表字段异常，实际字段为：{stock_info.columns.tolist()}")

    df["code"] = df["code"].apply(normalize_code)
    df["name"] = df["name"].astype(str).str.strip()
    df = df[df["code"].apply(is_mainboard_code)]
    df = df[~df["name"].apply(is_bad_name)]
    df = df.drop_duplicates(subset=["code"]).sort_values("code").reset_index(drop=True)

    print(f"AkShare 获取主板股票数量：{len(df)}")
    return df


def get_mainboard_stocks_from_local(data_dir: str = "data") -> pd.DataFrame:
    codes = scan_local_stock_codes(data_dir)

    if not codes:
        raise RuntimeError(f"{data_dir} 目录下没有找到主板股票 CSV 文件。")

    name_map = load_local_stock_names(data_dir)

    df = pd.DataFrame({
        "code": codes,
        "name": [name_map.get(c, "") for c in codes],
    })

    if name_map:
        before = len(df)
        df = df[~df["name"].apply(is_bad_name)].copy()
        after = len(df)

        if before != after:
            print(f"根据本地名称文件剔除 ST/退市股票：{before - after} 只")

    df = df.drop_duplicates(subset=["code"]).reset_index(drop=True)

    print(f"从本地 {data_dir} 目录识别主板股票数量：{len(df)}")
    return df


def get_mainboard_stocks_from_ak_and_local(data_dir: str = "data") -> pd.DataFrame:
    local_codes = set(scan_local_stock_codes(data_dir))

    if not local_codes:
        raise RuntimeError(f"{data_dir} 目录下没有找到主板股票 CSV 文件。")

    ak_df = get_a_stock_universe_from_ak()
    save_stock_names(data_dir, ak_df)

    df = ak_df[ak_df["code"].isin(local_codes)].copy()
    df = df.sort_values("code").reset_index(drop=True)

    print(f"本地 CSV 股票数量：{len(local_codes)}")
    print(f"AkShare 主板股票数量：{len(ak_df)}")
    print(f"两者交集数量：{len(df)}")

    if df.empty:
        raise RuntimeError("AkShare 股票列表和本地 CSV 没有交集，请检查代码格式。")

    return df


def get_universe(data_dir: str = "data", source: str = "local") -> pd.DataFrame:
    source = str(source).lower().strip()

    if source == "local":
        return get_mainboard_stocks_from_local(data_dir)

    if source == "ak":
        return get_mainboard_stocks_from_ak_and_local(data_dir)

    if source == "auto":
        try:
            return get_mainboard_stocks_from_ak_and_local(data_dir)
        except Exception as e:
            print(f"AkShare 获取股票列表失败，回退到纯本地模式。错误：{e}")
            return get_mainboard_stocks_from_local(data_dir)

    raise ValueError(f"未知 universe-source：{source}")


def build_code_name_map(universe: pd.DataFrame) -> Dict[str, str]:
    if universe is None or universe.empty:
        return {}

    if "code" not in universe.columns:
        return {}

    name_col = None

    for c in ["name", "名称", "stock_name", "股票简称"]:
        if c in universe.columns:
            name_col = c
            break

    if name_col is None:
        return {}

    tmp = universe[["code", name_col]].copy()
    tmp.columns = ["code", "name"]
    tmp["code"] = tmp["code"].astype(str).apply(normalize_code)
    tmp["name"] = tmp["name"].astype(str).str.strip()
    tmp = tmp[tmp["name"] != ""]
    tmp = tmp[tmp["name"].str.lower() != "nan"]

    return dict(zip(tmp["code"], tmp["name"]))


# ==============================
# 3. 指标计算
# ==============================

def add_indicators(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    df["pre_close"] = df["close"].shift(1)

    df["ma20"] = df["close"].rolling(20).mean()
    df["ma60"] = df["close"].rolling(60).mean()
    df["ma120"] = df["close"].rolling(120).mean()

    df["ma20_slope_5"] = df["ma20"] / df["ma20"].shift(5) - 1
    df["ma60_slope_10"] = df["ma60"] / df["ma60"].shift(10) - 1

    df["ret10"] = df["close"] / df["close"].shift(10) - 1
    df["ret20"] = df["close"] / df["close"].shift(20) - 1
    df["ret60"] = df["close"] / df["close"].shift(60) - 1

    df["amount_ma5"] = df["amount"].rolling(5).mean()
    df["amount_ma20"] = df["amount"].rolling(20).mean()
    df["amount_ratio_5_20"] = df["amount_ma5"] / df["amount_ma20"]

    df["turnover_ma5"] = df["turnover"].rolling(5).mean() if "turnover" in df.columns else np.nan
    df["turnover_ma20"] = df["turnover"].rolling(20).mean() if "turnover" in df.columns else np.nan

    df["high20"] = df["high"].rolling(20).max()
    df["high60"] = df["high"].rolling(60).max()
    df["is_20d_high"] = df["close"] >= df["high20"] * 0.999
    df["is_60d_high"] = df["close"] >= df["high60"] * 0.999
    df["dist_to_60d_high"] = df["close"] / df["high60"] - 1

    df["max_dd20"] = rolling_max_drawdown_np(df["close"], window=20)

    df["is_up_day"] = df["close"] > df["pre_close"]
    df["is_down_day"] = df["close"] < df["pre_close"]

    up_amount_20 = df["amount"].where(df["is_up_day"], 0.0).rolling(20, min_periods=20).sum()
    down_amount_20 = df["amount"].where(df["is_down_day"], 0.0).rolling(20, min_periods=20).sum()
    df["up_down_amount_ratio20"] = up_amount_20 / down_amount_20.replace(0, np.nan)

    df["is_limit_up"] = df["pct_chg"] >= 9.5
    df["limit_up_count_5"] = df["is_limit_up"].rolling(5).sum()

    intraday_range = df["high"] - df["low"]
    close_position = np.where(
        intraday_range > 0,
        (df["close"] - df["low"]) / intraday_range,
        np.nan,
    )

    df["heavy_bearish_candle"] = (
        (df["pct_chg"] <= -5)
        & (df["amount"] > df["amount_ma20"] * 1.5)
        & (close_position <= 0.35)
    )

    return df


def calc_score(row: pd.Series, bench_ret20: float, bench_ret60: float) -> int:
    score = 0

    if row["close"] > row["ma20"]:
        score += 5
    if row["ma20"] > row["ma60"]:
        score += 8
    if row["ma60"] > row["ma120"]:
        score += 7
    if row["ma20_slope_5"] > 0:
        score += 5
    if row["ma60_slope_10"] > 0:
        score += 5

    if row["ret20_rank_pct"] >= 0.70:
        score += 8
    if row["ret60_rank_pct"] >= 0.70:
        score += 8

    if not pd.isna(bench_ret20):
        if row["ret20"] > bench_ret20 + 0.05:
            score += 5
        elif row["ret20"] > bench_ret20:
            score += 3

    if not pd.isna(bench_ret60):
        if row["ret60"] > bench_ret60 + 0.10:
            score += 4
        elif row["ret60"] > bench_ret60:
            score += 2

    ratio = row["amount_ratio_5_20"]

    if 1.2 <= ratio <= 4:
        score += 6
    elif 1.0 <= ratio < 1.2:
        score += 3

    if row["up_down_amount_ratio20"] > 1.2:
        score += 4
    elif row["up_down_amount_ratio20"] > 1.0:
        score += 2

    if (row["is_20d_high"] or row["is_60d_high"]) and ratio >= 1.2:
        score += 3

    if ratio < 4:
        score += 2

    if row["is_20d_high"]:
        score += 4
    if row["is_60d_high"]:
        score += 6

    if row["dist_to_60d_high"] >= -0.03:
        score += 3
    elif row["dist_to_60d_high"] >= -0.05:
        score += 2

    if row["close"] > row["ma20"] and row["dist_to_60d_high"] >= -0.05:
        score += 2

    max_dd20 = row["max_dd20"]

    if not pd.isna(max_dd20):
        if max_dd20 > -0.10:
            score += 5
        elif max_dd20 > -0.15:
            score += 3
        elif max_dd20 > -0.20:
            score += 1

    overheat = False

    if not pd.isna(row["ret10"]) and row["ret10"] > 0.35:
        overheat = True
    if not pd.isna(row["ret20"]) and row["ret20"] > 0.60:
        overheat = True
    if not pd.isna(row["limit_up_count_5"]) and row["limit_up_count_5"] >= 3:
        overheat = True

    if not overheat:
        score += 4

    if not row["heavy_bearish_candle"]:
        score += 3

    turnover_ma5 = row["turnover_ma5"]

    if pd.isna(turnover_ma5):
        score += 2
    else:
        if turnover_ma5 < 15:
            score += 3
        elif turnover_ma5 < 25:
            score += 1

    return int(min(score, 100))


def is_overheat(row: pd.Series) -> bool:
    if not pd.isna(row["ret10"]) and row["ret10"] > 0.35:
        return True
    if not pd.isna(row["ret20"]) and row["ret20"] > 0.60:
        return True
    if not pd.isna(row["limit_up_count_5"]) and row["limit_up_count_5"] >= 3:
        return True
    return False


def classify_status(row: pd.Series) -> str:
    if not row["basic_liquid"]:
        return "流动性不足"

    if row["close"] < row["ma60"]:
        return "跌破MA60剔除"

    if row["close"] < row["ma20"]:
        return "跌破MA20观察"

    if row["heavy_bearish_candle"]:
        return "放量长阴观察"

    if is_overheat(row):
        return "短期过热"

    if row["score"] >= 85 and row["candidate"]:
        return "强趋势池"

    if row["score"] >= MIN_SCORE and row["candidate"]:
        return "趋势观察池"

    if (
        row["close"] > row["ma60"]
        and row["ma20"] > row["ma60"]
        and abs(row["close"] / row["ma20"] - 1) <= 0.05
        and row["ret60_rank_pct"] >= 0.60
    ):
        return "回踩观察"

    return "剔除"


# ==============================
# 4. 数据加载
# ==============================

def fetch_benchmark_data() -> pd.DataFrame:
    if not os.path.exists(BENCHMARK_CACHE_FILE):
        raise FileNotFoundError(
            f"基准缓存文件不存在：{BENCHMARK_CACHE_FILE}\n"
            f"请先准备好沪深300缓存文件，并放入 data 目录。"
        )

    try:
        df = pd.read_csv(BENCHMARK_CACHE_FILE)
    except Exception as e:
        raise RuntimeError(f"读取基准文件失败：{BENCHMARK_CACHE_FILE}，错误：{e}")

    df = standardize_local_df(df, code="沪深300")

    if df is None or df.empty:
        raise ValueError(f"基准缓存文件为空或字段异常：{BENCHMARK_CACHE_FILE}")

    return df


def load_stock_history_from_local(
    code: str,
    data_dir: str = "data",
    min_bars: int = 130,
) -> Optional[pd.DataFrame]:
    code = normalize_code(code)
    cache_path = os.path.join(data_dir, f"{code}.csv")

    if not os.path.exists(cache_path):
        return None

    try:
        df = pd.read_csv(cache_path)
        df = standardize_local_df(df, code=code)

        if df is None or df.empty:
            return None

        if len(df) < min_bars:
            return None

        return df

    except Exception as e:
        print(f"{code} 读取本地数据失败：{e}")
        return None


def preload_all_stock_data(codes: List[str]) -> Dict[str, pd.DataFrame]:
    data = {}

    print("正在从本地 data 目录读取股票历史数据...")
    print(f"数据目录: {CACHE_DIR}")
    print("注意：本步骤不会下载股票历史行情。")

    missing_count = 0
    insufficient_count = 0
    invalid_count = 0

    for code in tqdm(codes, desc="股票数据读取"):
        code = normalize_code(code)

        try:
            cache_path = os.path.join(CACHE_DIR, f"{code}.csv")

            if not os.path.exists(cache_path):
                missing_count += 1
                continue

            df = load_stock_history_from_local(
                code=code,
                data_dir=CACHE_DIR,
                min_bars=MIN_BARS,
            )

            if df is None or df.empty:
                insufficient_count += 1
                continue

            required_price_cols = ["open", "high", "low", "close", "amount"]
            missing_cols = [c for c in required_price_cols if c not in df.columns]

            if missing_cols:
                invalid_count += 1
                print(f"{code} 缺少字段 {missing_cols}，跳过。")
                continue

            df = add_indicators(df)
            data[code] = df

        except Exception as e:
            invalid_count += 1
            print(f"{code} 读取或处理失败：{e}")
            continue

    print(f"成功加载 {len(data)} 只股票数据。")
    print(f"本地文件不存在数量: {missing_count}")
    print(f"数据不足或为空数量: {insufficient_count}")
    print(f"字段异常或处理失败数量: {invalid_count}")

    return data


# ==============================
# 5. 指定日期股票池构建
# ==============================

def resolve_signal_date(
    requested_date: str,
    bench_df: pd.DataFrame,
    use_prev_trading_day: bool = False,
) -> pd.Timestamp:
    dt = pd.Timestamp(requested_date).normalize()
    trading_dates = bench_df["date"].drop_duplicates().sort_values().reset_index(drop=True)

    if (trading_dates == dt).any():
        return dt

    if not use_prev_trading_day:
        raise ValueError(
            f"{requested_date} 不是基准交易日。"
            f"如果想自动使用前一个交易日，请加参数 --use-prev-trading-day"
        )

    prev_dates = trading_dates[trading_dates < dt]

    if prev_dates.empty:
        raise ValueError(f"{requested_date} 之前没有可用交易日。")

    resolved = prev_dates.iloc[-1]
    print(f"输入日期 {requested_date} 不是交易日，已自动使用前一个交易日：{resolved.date()}")

    return resolved


def build_pool_on_date(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Optional[Dict[str, str]] = None,
) -> pd.DataFrame:
    if code_name_map is None:
        code_name_map = {}

    signal_date = pd.Timestamp(signal_date).normalize()

    bench_slice = bench_df[bench_df["date"] <= signal_date].copy()

    if len(bench_slice) < 61:
        return pd.DataFrame()

    bench_close_date = bench_slice.iloc[-1]["close"]
    bench_close_20 = bench_slice.iloc[-21]["close"]
    bench_close_60 = bench_slice.iloc[-61]["close"]

    bench_ret20 = bench_close_date / bench_close_20 - 1
    bench_ret60 = bench_close_date / bench_close_60 - 1

    required_indicator_cols = [
        "close", "ma20", "ma60", "ma120",
        "ma20_slope_5", "ma60_slope_10",
        "ret20", "ret60",
        "amount_ma20", "amount_ratio_5_20",
        "high20", "high60", "dist_to_60d_high",
        "max_dd20", "up_down_amount_ratio20",
    ]

    rows = []

    for code, df in stock_data.items():
        if df is None or df.empty:
            continue

        df_slice = df[df["date"] <= signal_date]

        if len(df_slice) < MIN_BARS:
            continue

        latest = df_slice.iloc[-1]

        if pd.Timestamp(latest["date"]).normalize() != signal_date:
            continue

        has_nan = False

        for col in required_indicator_cols:
            if col not in latest.index or pd.isna(latest[col]):
                has_nan = True
                break

        if has_nan:
            continue

        rows.append({
            "code": str(code),
            "name": code_name_map.get(str(code), ""),
            "signal_date": signal_date.strftime("%Y-%m-%d"),

            "close": latest["close"],
            "ma20": latest["ma20"],
            "ma60": latest["ma60"],
            "ma120": latest["ma120"],
            "ma20_slope_5": latest["ma20_slope_5"],
            "ma60_slope_10": latest["ma60_slope_10"],

            "ret10": latest["ret10"],
            "ret20": latest["ret20"],
            "ret60": latest["ret60"],

            "amount_ma5": latest["amount_ma5"],
            "amount_ma20": latest["amount_ma20"],
            "amount_ratio_5_20": latest["amount_ratio_5_20"],

            "turnover_ma5": latest.get("turnover_ma5", np.nan),
            "turnover_ma20": latest.get("turnover_ma20", np.nan),

            "high20": latest["high20"],
            "high60": latest["high60"],
            "is_20d_high": bool(latest["is_20d_high"]),
            "is_60d_high": bool(latest["is_60d_high"]),
            "dist_to_60d_high": latest["dist_to_60d_high"],

            "max_dd20": latest["max_dd20"],
            "up_down_amount_ratio20": latest["up_down_amount_ratio20"],

            "limit_up_count_5": latest.get("limit_up_count_5", np.nan),
            "heavy_bearish_candle": bool(latest.get("heavy_bearish_candle", False)),

            "bench_ret20": bench_ret20,
            "bench_ret60": bench_ret60,
        })

    if not rows:
        return pd.DataFrame()

    df_pool = pd.DataFrame(rows)

    df_pool["ret20_rank_pct"] = df_pool["ret20"].rank(pct=True)
    df_pool["ret60_rank_pct"] = df_pool["ret60"].rank(pct=True)

    df_pool["basic_liquid"] = df_pool["amount_ma20"] >= MIN_AVG_AMOUNT_20

    df_pool["trend_basic"] = (
        (df_pool["close"] > df_pool["ma20"])
        & (df_pool["ma20"] > df_pool["ma60"])
        & (df_pool["ma60"] > df_pool["ma120"])
        & (df_pool["ma20_slope_5"] > 0)
        & (df_pool["ma60_slope_10"] > 0)
    )

    df_pool["relative_strength"] = (
        (df_pool["ret20_rank_pct"] >= 0.70)
        & (df_pool["ret60_rank_pct"] >= 0.70)
    )

    df_pool["volume_ok"] = (
        (df_pool["amount_ratio_5_20"] >= 1.2)
        & (df_pool["amount_ratio_5_20"] <= 4)
    )

    df_pool["near_breakout"] = (
        (df_pool["dist_to_60d_high"] >= -0.05)
        | df_pool["is_20d_high"]
        | df_pool["is_60d_high"]
    )

    df_pool["risk_ok"] = (
        (df_pool["close"] > df_pool["ma20"])
        & (df_pool["max_dd20"] > -0.20)
        & (~df_pool["heavy_bearish_candle"])
    )

    df_pool["candidate"] = (
        df_pool["basic_liquid"]
        & df_pool["trend_basic"]
        & df_pool["relative_strength"]
        & df_pool["volume_ok"]
        & df_pool["near_breakout"]
        & df_pool["risk_ok"]
    )

    df_pool["score"] = df_pool.apply(
        lambda row: calc_score(row, bench_ret20, bench_ret60),
        axis=1,
    )

    df_pool["status"] = df_pool.apply(classify_status, axis=1)

    status_order = {
        "强趋势池": 1,
        "趋势观察池": 2,
        "回踩观察": 3,
        "跌破MA20观察": 4,
        "放量长阴观察": 5,
        "短期过热": 6,
        "跌破MA60剔除": 7,
        "流动性不足": 8,
        "剔除": 9,
    }

    df_pool["_order"] = df_pool["status"].map(status_order).fillna(99)
    df_pool = df_pool.sort_values(
        ["_order", "score"],
        ascending=[True, False],
    ).drop(columns=["_order"]).reset_index(drop=True)

    return df_pool


# ==============================
# 6. 未来走势与动态卖出计算
# ==============================

def get_forward_trading_dates(
    signal_date: pd.Timestamp,
    bench_df: pd.DataFrame,
    forward_days: int,
) -> List[pd.Timestamp]:
    signal_date = pd.Timestamp(signal_date).normalize()

    future_dates = bench_df.loc[
        bench_df["date"] > signal_date,
        "date"
    ].head(forward_days).tolist()

    return [pd.Timestamp(d).normalize() for d in future_dates]


def calc_forward_return_path_by_calendar(
    df: pd.DataFrame,
    forward_dates: List[pd.Timestamp],
    buy_price_field: str = BUY_PRICE_FIELD,
) -> Dict:
    if df is None or df.empty or not forward_dates:
        return {}

    df_map = df.set_index("date", drop=False)
    buy_date = forward_dates[0]

    if buy_date not in df_map.index:
        return {}

    buy_row = df_map.loc[buy_date]

    if isinstance(buy_row, pd.DataFrame):
        buy_row = buy_row.iloc[-1]

    buy_price = get_row_price(buy_row, buy_price_field, fallback="close")

    if pd.isna(buy_price) or buy_price <= 0:
        return {}

    result = {
        "buy_date": buy_date.strftime("%Y-%m-%d"),
        "buy_price": buy_price,
    }

    returns = []

    for i, d in enumerate(forward_dates, start=1):
        if d not in df_map.index:
            ret = np.nan
        else:
            row = df_map.loc[d]

            if isinstance(row, pd.DataFrame):
                row = row.iloc[-1]

            close_price = get_row_price(row, "close", fallback="close")

            if pd.isna(close_price) or close_price <= 0:
                ret = np.nan
            else:
                ret = close_price / buy_price - 1

        result[f"D{i}"] = ret
        returns.append(ret)

    ret_series = pd.Series(returns).dropna()
    n = len(forward_dates)

    if len(ret_series) == 0:
        result[f"max_ret_{n}"] = np.nan
        result[f"min_ret_{n}"] = np.nan
        result[f"max_dd_{n}"] = np.nan
        result["day_to_max_ret"] = np.nan
        result["first_profit_5_day"] = np.nan
        result["first_profit_10_day"] = np.nan
        result["first_drawdown_5_day"] = np.nan
        return result

    result[f"max_ret_{n}"] = ret_series.max()
    result[f"min_ret_{n}"] = ret_series.min()
    result[f"max_dd_{n}"] = max_drawdown(1 + ret_series)
    result["day_to_max_ret"] = int(ret_series.idxmax() + 1)

    first_profit_5 = np.nan
    first_profit_10 = np.nan
    first_drawdown_5 = np.nan

    for idx, ret in enumerate(returns, start=1):
        if pd.isna(ret):
            continue
        if pd.isna(first_profit_5) and ret >= 0.05:
            first_profit_5 = idx
        if pd.isna(first_profit_10) and ret >= 0.10:
            first_profit_10 = idx
        if pd.isna(first_drawdown_5) and ret <= -0.05:
            first_drawdown_5 = idx

    result["first_profit_5_day"] = first_profit_5
    result["first_profit_10_day"] = first_profit_10
    result["first_drawdown_5_day"] = first_drawdown_5

    return result


def judge_weakness_for_exit(
    row: pd.Series,
    buy_price: float,
    peak_close: float,
    hold_day: int,
    min_hold_days: int = 3,
    stop_loss_pct: float = -0.08,
    trailing_start_pct: float = 0.10,
    trailing_drawdown_pct: float = 0.08,
) -> Optional[str]:
    close = row.get("close", np.nan)
    ma20 = row.get("ma20", np.nan)
    ma60 = row.get("ma60", np.nan)
    ma20_slope_5 = row.get("ma20_slope_5", np.nan)
    ret10 = row.get("ret10", np.nan)
    heavy_bearish = bool(row.get("heavy_bearish_candle", False))

    if pd.isna(close) or close <= 0 or pd.isna(buy_price) or buy_price <= 0:
        return None

    current_ret = close / buy_price - 1

    if current_ret <= stop_loss_pct:
        return f"硬止损{stop_loss_pct:.0%}"

    if not pd.isna(ma60) and close < ma60:
        return "跌破MA60"

    if hold_day < min_hold_days:
        return None

    if not pd.isna(peak_close) and peak_close > 0:
        peak_ret = peak_close / buy_price - 1
        drawdown_from_peak = close / peak_close - 1

        if peak_ret >= trailing_start_pct and drawdown_from_peak <= -trailing_drawdown_pct:
            return f"移动止盈回撤{trailing_drawdown_pct:.0%}"

    if heavy_bearish and not pd.isna(ma20) and close < ma20:
        return "放量长阴跌破MA20"

    if (
        not pd.isna(ma20)
        and not pd.isna(ma20_slope_5)
        and close < ma20
        and ma20_slope_5 <= 0
    ):
        return "跌破MA20且均线走弱"

    if (
        not pd.isna(ma20)
        and not pd.isna(ret10)
        and close < ma20
        and ret10 <= -0.08
    ):
        return "10日转弱且跌破MA20"

    return None


def calc_forward_trade_with_dynamic_exit(
    df: pd.DataFrame,
    forward_dates: List[pd.Timestamp],
    buy_price_field: str = BUY_PRICE_FIELD,
    min_hold_days: int = 3,
    exit_sell_timing: str = "next_open",
    stop_loss_pct: float = -0.08,
    trailing_start_pct: float = 0.10,
    trailing_drawdown_pct: float = 0.08,
) -> Dict:
    if df is None or df.empty or not forward_dates:
        return {}

    exit_sell_timing = str(exit_sell_timing).lower().strip()

    if exit_sell_timing not in ["close", "next_open"]:
        raise ValueError("exit_sell_timing 只能是 close 或 next_open")

    df_map = df.set_index("date", drop=False)
    buy_date = forward_dates[0]

    if buy_date not in df_map.index:
        return {}

    buy_row = df_map.loc[buy_date]

    if isinstance(buy_row, pd.DataFrame):
        buy_row = buy_row.iloc[-1]

    buy_price = get_row_price(buy_row, buy_price_field, fallback="close")

    if pd.isna(buy_price) or buy_price <= 0:
        return {}

    result = {
        "dynamic_buy_date": buy_date.strftime("%Y-%m-%d"),
        "dynamic_buy_price": buy_price,
    }

    exited = False
    pending_exit = False
    pending_reason = None

    sell_date = None
    sell_price = np.nan
    sell_day = np.nan
    exit_reason = None
    final_ret = np.nan

    peak_close = np.nan
    last_valid_close = np.nan
    last_valid_day = np.nan
    last_valid_date = None

    strategy_returns = []

    for i, d in enumerate(forward_dates, start=1):
        strategy_ret = np.nan

        if d not in df_map.index:
            if exited and not pd.isna(final_ret):
                strategy_ret = final_ret

            result[f"strategy_D{i}"] = strategy_ret
            strategy_returns.append(strategy_ret)
            continue

        row = df_map.loc[d]

        if isinstance(row, pd.DataFrame):
            row = row.iloc[-1]

        close_price = get_row_price(row, "close", fallback="close")
        open_price = get_row_price(row, "open", fallback="close")

        if not pd.isna(close_price) and close_price > 0:
            last_valid_close = close_price
            last_valid_day = i
            last_valid_date = d

            if pd.isna(peak_close):
                peak_close = close_price
            else:
                peak_close = max(peak_close, close_price)

        if (not exited) and pending_exit:
            if pd.isna(open_price) or open_price <= 0:
                open_price = close_price

            if not pd.isna(open_price) and open_price > 0:
                sell_date = d
                sell_price = open_price
                sell_day = i
                exit_reason = pending_reason
                final_ret = sell_price / buy_price - 1
                strategy_ret = final_ret
                exited = True

            result[f"strategy_D{i}"] = strategy_ret
            strategy_returns.append(strategy_ret)
            continue

        if exited:
            strategy_ret = final_ret
            result[f"strategy_D{i}"] = strategy_ret
            strategy_returns.append(strategy_ret)
            continue

        if not pd.isna(close_price) and close_price > 0:
            strategy_ret = close_price / buy_price - 1

        result[f"strategy_D{i}"] = strategy_ret
        strategy_returns.append(strategy_ret)

        if pd.isna(close_price) or close_price <= 0:
            continue

        reason = judge_weakness_for_exit(
            row=row,
            buy_price=buy_price,
            peak_close=peak_close,
            hold_day=i,
            min_hold_days=min_hold_days,
            stop_loss_pct=stop_loss_pct,
            trailing_start_pct=trailing_start_pct,
            trailing_drawdown_pct=trailing_drawdown_pct,
        )

        if reason is None:
            continue

        if exit_sell_timing == "close":
            sell_date = d
            sell_price = close_price
            sell_day = i
            exit_reason = reason
            final_ret = sell_price / buy_price - 1
            result[f"strategy_D{i}"] = final_ret
            strategy_returns[-1] = final_ret
            exited = True

        elif exit_sell_timing == "next_open":
            if i >= len(forward_dates):
                sell_date = d
                sell_price = close_price
                sell_day = i
                exit_reason = reason + "；末日收盘卖出"
                final_ret = sell_price / buy_price - 1
                result[f"strategy_D{i}"] = final_ret
                strategy_returns[-1] = final_ret
                exited = True
            else:
                pending_exit = True
                pending_reason = reason

    if not exited:
        if not pd.isna(last_valid_close) and last_valid_close > 0:
            sell_date = last_valid_date
            sell_price = last_valid_close
            sell_day = last_valid_day
            exit_reason = "持有到期"
            final_ret = sell_price / buy_price - 1
        else:
            sell_date = np.nan
            sell_price = np.nan
            sell_day = np.nan
            exit_reason = "无有效卖出价"
            final_ret = np.nan

    ret_series = pd.Series(strategy_returns).dropna()
    n = len(forward_dates)

    result["dynamic_sell_date"] = (
        sell_date.strftime("%Y-%m-%d") if isinstance(sell_date, pd.Timestamp) else sell_date
    )
    result["dynamic_sell_price"] = sell_price
    result["dynamic_sell_day"] = sell_day
    result["dynamic_exit_reason"] = exit_reason
    result["dynamic_final_ret"] = final_ret

    if len(ret_series) == 0:
        result[f"dynamic_max_ret_{n}"] = np.nan
        result[f"dynamic_min_ret_{n}"] = np.nan
        result[f"dynamic_max_dd_{n}"] = np.nan
    else:
        result[f"dynamic_max_ret_{n}"] = ret_series.max()
        result[f"dynamic_min_ret_{n}"] = ret_series.min()
        result[f"dynamic_max_dd_{n}"] = max_drawdown(1 + ret_series)

    return result


# ==============================
# 7. 统计表
# ==============================

def build_excess_detail(
    detail_df: pd.DataFrame,
    forward_days: int,
) -> pd.DataFrame:
    if detail_df is None or detail_df.empty:
        return pd.DataFrame()

    bench_rows = detail_df[detail_df["row_type"] == "benchmark"]
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    if bench_rows.empty or stock_rows.empty:
        return pd.DataFrame()

    bench_row = bench_rows.iloc[0]
    rows = []

    for _, row in stock_rows.iterrows():
        item = {
            "code": row["code"],
            "name": row.get("name", ""),
            "status": row["status"],
            "score": row["score"],
            "signal_date": row["signal_date"],
            "buy_date": row["buy_date"],
        }

        outperform_days = 0
        valid_days = 0
        excess_values = []

        for h in range(1, forward_days + 1):
            col = f"D{h}"
            stock_ret = row.get(col, np.nan)
            bench_ret = bench_row.get(col, np.nan)

            if pd.isna(stock_ret) or pd.isna(bench_ret):
                excess = np.nan
            else:
                excess = stock_ret - bench_ret
                valid_days += 1
                excess_values.append(excess)

                if excess > 0:
                    outperform_days += 1

            item[f"excess_D{h}"] = excess

        item["avg_excess"] = np.nan if len(excess_values) == 0 else np.mean(excess_values)
        item["outperform_days"] = outperform_days
        item["valid_days"] = valid_days

        rows.append(item)

    return pd.DataFrame(rows)


def build_forward_summary(
    detail_df: pd.DataFrame,
    forward_days: int,
    group_col: str = "status",
) -> pd.DataFrame:
    if detail_df is None or detail_df.empty:
        return pd.DataFrame()

    bench_rows = detail_df[detail_df["row_type"] == "benchmark"]
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    if bench_rows.empty or stock_rows.empty:
        return pd.DataFrame()

    bench_row = bench_rows.iloc[0]
    rows = []

    for group_value, group in stock_rows.groupby(group_col):
        for h in range(1, forward_days + 1):
            col = f"D{h}"

            if col not in group.columns:
                continue

            bench_ret = bench_row[col]
            valid = group[group[col].notna()].copy()

            if valid.empty:
                continue

            ret = valid[col]
            excess = ret - bench_ret

            rows.append({
                group_col: group_value,
                "horizon": h,
                "count": len(valid),
                "avg_ret": ret.mean(),
                "median_ret": ret.median(),
                "win_rate": (ret > 0).mean(),
                "bench_ret": bench_ret,
                "avg_excess": excess.mean(),
                "outperform_rate": (excess > 0).mean(),
            })

    return pd.DataFrame(rows)


def build_overall_forward_summary(
    detail_df: pd.DataFrame,
    forward_days: int,
) -> pd.DataFrame:
    if detail_df is None or detail_df.empty:
        return pd.DataFrame()

    bench_rows = detail_df[detail_df["row_type"] == "benchmark"]
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    if bench_rows.empty or stock_rows.empty:
        return pd.DataFrame()

    bench_row = bench_rows.iloc[0]
    rows = []

    for h in range(1, forward_days + 1):
        col = f"D{h}"

        if col not in stock_rows.columns:
            continue

        bench_ret = bench_row[col]
        valid = stock_rows[stock_rows[col].notna()].copy()

        if valid.empty:
            continue

        ret = valid[col]
        excess = ret - bench_ret

        rows.append({
            "horizon": h,
            "count": len(valid),
            "avg_ret": ret.mean(),
            "median_ret": ret.median(),
            "win_rate": (ret > 0).mean(),
            "bench_ret": bench_ret,
            "avg_excess": excess.mean(),
            "outperform_rate": (excess > 0).mean(),
        })

    return pd.DataFrame(rows)


def build_dynamic_exit_summary(
    dynamic_exit_df: pd.DataFrame,
    group_col: str = "status",
) -> pd.DataFrame:
    if dynamic_exit_df is None or dynamic_exit_df.empty:
        return pd.DataFrame()

    rows = []

    def summarize_one(group_name, group):
        ret = group["dynamic_final_ret"].dropna()

        if ret.empty:
            return

        max_ret_cols = group.filter(like="dynamic_max_ret_").columns
        max_dd_cols = group.filter(like="dynamic_max_dd_").columns

        rows.append({
            "group": group_name,
            "count": len(ret),
            "avg_final_ret": ret.mean(),
            "median_final_ret": ret.median(),
            "win_rate": (ret > 0).mean(),
            "avg_hold_days": group["dynamic_sell_day"].dropna().mean()
            if "dynamic_sell_day" in group.columns else np.nan,
            "avg_max_ret": group[max_ret_cols[0]].mean()
            if len(max_ret_cols) > 0 else np.nan,
            "avg_max_dd": group[max_dd_cols[0]].mean()
            if len(max_dd_cols) > 0 else np.nan,
        })

    summarize_one("全部", dynamic_exit_df)

    if group_col in dynamic_exit_df.columns:
        for group_name, group in dynamic_exit_df.groupby(group_col):
            summarize_one(group_name, group)

    return pd.DataFrame(rows)


def build_dynamic_exit_reason_summary(dynamic_exit_df: pd.DataFrame) -> pd.DataFrame:
    if dynamic_exit_df is None or dynamic_exit_df.empty:
        return pd.DataFrame()

    if "dynamic_exit_reason" not in dynamic_exit_df.columns:
        return pd.DataFrame()

    out = (
        dynamic_exit_df
        .groupby(["status", "dynamic_exit_reason"])
        .size()
        .reset_index(name="count")
        .sort_values(["status", "count"], ascending=[True, False])
        .reset_index(drop=True)
    )

    out = out.rename(columns={
        "status": "分组",
        "dynamic_exit_reason": "卖出原因",
        "count": "数量",
    })

    return out


# ==============================
# 8. Excel 美化输出
# ==============================

def get_key_horizons(forward_days: int) -> List[int]:
    base = [1, 3, 5, 10, 20, 30, 60]
    return [h for h in base if h <= forward_days]


def style_worksheet(
    ws,
    freeze: str = "A2",
    percent_keywords: Optional[List[str]] = None,
    money_keywords: Optional[List[str]] = None,
    integer_keywords: Optional[List[str]] = None,
    apply_red_green: bool = True,
):
    if percent_keywords is None:
        percent_keywords = [
            "收益", "超额", "涨幅", "回撤", "胜率", "比例", "D",
            "平均超额", "沪深300收益", "距60日高点", "硬止损", "移动止盈",
        ]

    if money_keywords is None:
        money_keywords = ["成交额", "买入价", "卖出价", "收盘"]

    if integer_keywords is None:
        integer_keywords = ["天数", "样本数", "评分", "跑赢天数", "有效天数", "持有天数", "数量"]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    max_row = ws.max_row
    max_col = ws.max_column

    positive_fill = PatternFill("solid", fgColor="FCE4D6")
    negative_fill = PatternFill("solid", fgColor="E2F0D9")

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        header = ws.cell(row=1, column=col_idx).value
        header_str = "" if header is None else str(header)

        max_len = len(header_str)

        for row_idx in range(2, min(max_row, 200) + 1):
            v = ws.cell(row=row_idx, column=col_idx).value

            if v is not None:
                max_len = max(max_len, len(str(v)))

        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 28)

        is_percent_col = any(k in header_str for k in percent_keywords)
        is_money_col = any(k in header_str for k in money_keywords)
        is_integer_col = any(k in header_str for k in integer_keywords)

        for row_idx in range(2, max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = Alignment(vertical="center")

            if is_percent_col and isinstance(cell.value, (int, float)) and not pd.isna(cell.value):
                cell.number_format = "0.00%"
            elif is_money_col and isinstance(cell.value, (int, float)) and not pd.isna(cell.value):
                cell.number_format = "#,##0.00"
            elif is_integer_col and isinstance(cell.value, (int, float)) and not pd.isna(cell.value):
                cell.number_format = "0"

        if apply_red_green and is_percent_col and max_row >= 2:
            data_range = f"{col_letter}2:{col_letter}{max_row}"

            ws.conditional_formatting.add(
                data_range,
                CellIsRule(
                    operator="greaterThan",
                    formula=["0"],
                    fill=positive_fill,
                ),
            )

            ws.conditional_formatting.add(
                data_range,
                CellIsRule(
                    operator="lessThan",
                    formula=["0"],
                    fill=negative_fill,
                ),
            )


# ==============================
# 8.1 新增：指定日期分类模块
# ==============================

def build_classification_summary(pool_df: pd.DataFrame) -> pd.DataFrame:
    if pool_df is None or pool_df.empty:
        return pd.DataFrame()

    total = len(pool_df)

    summary = (
        pool_df
        .groupby("status")
        .agg(
            数量=("code", "count"),
            平均评分=("score", "mean"),
            最高评分=("score", "max"),
            平均20日涨幅=("ret20", "mean"),
            平均60日涨幅=("ret60", "mean"),
            平均20日成交额=("amount_ma20", "mean"),
            平均量能比=("amount_ratio_5_20", "mean"),
            平均20日最大回撤=("max_dd20", "mean"),
        )
        .reset_index()
        .rename(columns={"status": "分类"})
    )

    summary["占比"] = summary["数量"] / total

    status_order = {
        "强趋势池": 1,
        "趋势观察池": 2,
        "回踩观察": 3,
        "跌破MA20观察": 4,
        "放量长阴观察": 5,
        "短期过热": 6,
        "跌破MA60剔除": 7,
        "流动性不足": 8,
        "剔除": 9,
    }

    summary["排序"] = summary["分类"].map(status_order).fillna(99)
    summary = summary.sort_values(["排序", "数量"], ascending=[True, False])
    summary = summary.drop(columns=["排序"]).reset_index(drop=True)

    return summary


def make_classification_readable(pool_df: pd.DataFrame) -> pd.DataFrame:
    if pool_df is None or pool_df.empty:
        return pd.DataFrame()

    df = pool_df.copy()

    keep_cols = [
        "code", "name", "signal_date", "status", "score",
        "close", "ma20", "ma60", "ma120",
        "ma20_slope_5", "ma60_slope_10",
        "ret10", "ret20", "ret60",
        "ret20_rank_pct", "ret60_rank_pct",
        "amount_ma5", "amount_ma20", "amount_ratio_5_20",
        "turnover_ma5", "turnover_ma20",
        "is_20d_high", "is_60d_high", "dist_to_60d_high",
        "max_dd20", "up_down_amount_ratio20",
        "limit_up_count_5", "heavy_bearish_candle",
        "bench_ret20", "bench_ret60",
        "basic_liquid", "trend_basic", "relative_strength",
        "volume_ok", "near_breakout", "risk_ok", "candidate",
    ]

    keep_cols = [c for c in keep_cols if c in df.columns]
    out = df[keep_cols].copy()

    rename_map = {
        "code": "代码",
        "name": "名称",
        "signal_date": "信号日",
        "status": "分类",
        "score": "评分",
        "close": "收盘价",
        "ma20": "MA20",
        "ma60": "MA60",
        "ma120": "MA120",
        "ma20_slope_5": "MA20近5日斜率",
        "ma60_slope_10": "MA60近10日斜率",
        "ret10": "近10日涨幅",
        "ret20": "近20日涨幅",
        "ret60": "近60日涨幅",
        "ret20_rank_pct": "20日强度排名百分位",
        "ret60_rank_pct": "60日强度排名百分位",
        "amount_ma5": "5日均成交额",
        "amount_ma20": "20日均成交额",
        "amount_ratio_5_20": "量能比5/20",
        "turnover_ma5": "5日均换手率",
        "turnover_ma20": "20日均换手率",
        "is_20d_high": "是否接近20日新高",
        "is_60d_high": "是否接近60日新高",
        "dist_to_60d_high": "距60日高点",
        "max_dd20": "近20日最大回撤",
        "up_down_amount_ratio20": "近20日上涨/下跌成交额比",
        "limit_up_count_5": "近5日涨停次数",
        "heavy_bearish_candle": "是否放量长阴",
        "bench_ret20": "沪深300近20日涨幅",
        "bench_ret60": "沪深300近60日涨幅",
        "basic_liquid": "流动性达标",
        "trend_basic": "趋势结构达标",
        "relative_strength": "相对强度达标",
        "volume_ok": "量能达标",
        "near_breakout": "接近突破",
        "risk_ok": "风险过滤通过",
        "candidate": "是否候选",
    }

    out = out.rename(columns=rename_map)

    status_order = {
        "强趋势池": 1,
        "趋势观察池": 2,
        "回踩观察": 3,
        "跌破MA20观察": 4,
        "放量长阴观察": 5,
        "短期过热": 6,
        "跌破MA60剔除": 7,
        "流动性不足": 8,
        "剔除": 9,
    }

    if "分类" in out.columns:
        out["_排序"] = out["分类"].map(status_order).fillna(99)

        if "评分" in out.columns:
            out = out.sort_values(["_排序", "评分"], ascending=[True, False])
        else:
            out = out.sort_values(["_排序"], ascending=True)

        out = out.drop(columns=["_排序"]).reset_index(drop=True)

    return out


def save_classification_excel(
    output_xlsx: str,
    signal_date: pd.Timestamp,
    pool_df: pd.DataFrame,
) -> None:
    readable_df = make_classification_readable(pool_df)
    summary_df = build_classification_summary(pool_df)

    dashboard_rows = []
    total_count = len(pool_df) if pool_df is not None else 0

    dashboard_rows.append(["信号日期", signal_date.strftime("%Y-%m-%d"), "分类所使用的交易日"])
    dashboard_rows.append(["全市场有效样本数", total_count, "满足数据长度和指标完整性的股票数量"])
    dashboard_rows.append(["最低评分参数", MIN_SCORE, "趋势观察池最低评分"])
    dashboard_rows.append(["20日均成交额门槛", MIN_AVG_AMOUNT_20, "流动性过滤条件"])
    dashboard_rows.append(["最少K线数量", MIN_BARS, "参与计算要求的最低历史K线数量"])

    if summary_df is not None and not summary_df.empty:
        for _, r in summary_df.iterrows():
            dashboard_rows.append([
                f"{r['分类']}数量",
                r["数量"],
                f"占比 {r['占比']:.2%}",
            ])

    dashboard_df = pd.DataFrame(dashboard_rows, columns=["指标", "数值", "说明"])

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        dashboard_df.to_excel(writer, sheet_name="看板", index=False)
        summary_df.to_excel(writer, sheet_name="分类统计", index=False)
        readable_df.to_excel(writer, sheet_name="全市场分类", index=False)

        important_statuses = [
            "强趋势池",
            "趋势观察池",
            "回踩观察",
            "跌破MA20观察",
            "放量长阴观察",
            "短期过热",
            "跌破MA60剔除",
            "流动性不足",
            "剔除",
        ]

        if not readable_df.empty and "分类" in readable_df.columns:
            for status in important_statuses:
                sub = readable_df[readable_df["分类"] == status].copy()

                if sub.empty:
                    continue

                sheet_name = status[:31]
                sub.to_excel(writer, sheet_name=sheet_name, index=False)

        wb = writer.book

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            if sheet_name == "看板":
                style_worksheet(ws, freeze="A2", apply_red_green=True)
                ws.column_dimensions["A"].width = 28
                ws.column_dimensions["B"].width = 20
                ws.column_dimensions["C"].width = 56
            else:
                style_worksheet(ws, freeze="A2", apply_red_green=True)

def analyze_pool_classification(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Dict[str, str],
) -> None:
    signal_date = pd.Timestamp(signal_date).normalize()

    print("\n========== 指定日期股票分类 ==========")
    print(f"信号日期: {signal_date.date()}")
    print("功能: 仅分类，不计算未来收益")
    print("====================================")

    pool_df = build_pool_on_date(
        signal_date=signal_date,
        stock_data=stock_data,
        bench_df=bench_df,
        code_name_map=code_name_map,
    )

    if pool_df is None or pool_df.empty:
        print("指定日期未能构建股票分类结果。可能原因：")
        print("1. 该日期不是交易日；")
        print("2. 本地数据不足；")
        print("3. 当天股票没有交易数据；")
        print("4. 指标窗口不足。")
        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    output_xlsx = os.path.join(
        OUTPUT_DIR,
        f"{CLASSIFICATION_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.xlsx",
    )

    output_csv = os.path.join(
        OUTPUT_DIR,
        f"{CLASSIFICATION_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.csv",
    )

    # === Excel 保存 ===
    try:
        save_classification_excel(
            output_xlsx=output_xlsx,
            signal_date=signal_date,
            pool_df=pool_df,
        )
        print(f"\n分类结果 Excel 已保存到：{output_xlsx}")
    except Exception as e:
        print(f"保存分类 Excel 失败：{e}")

    # === CSV 保存（独立 try，确保无论 Excel 是否成功都会执行） ===
    try:
        pool_df.to_csv(output_csv, index=False, encoding="utf-8-sig")
        print(f"分类结果 CSV 已保存到：{output_csv}")
    except Exception as e:
        print(f"保存分类 CSV 失败：{e}")

    summary_df = build_classification_summary(pool_df)

    if not summary_df.empty:
        print("\n========== 分类统计 ==========")
        print(summary_df.to_string(index=False, formatters={
            "占比": "{:.2%}".format,
            "平均评分": "{:.2f}".format,
            "最高评分": "{:.2f}".format,
            "平均20日涨幅": "{:.2%}".format,
            "平均60日涨幅": "{:.2%}".format,
            "平均20日成交额": "{:,.0f}".format,
            "平均量能比": "{:.2f}".format,
            "平均20日最大回撤": "{:.2%}".format,
        }))

    selected_df = pool_df[pool_df["status"].isin(["强趋势池", "趋势观察池"])].copy()
    selected_df = selected_df.sort_values("score", ascending=False).reset_index(drop=True)

    if not selected_df.empty:
        preview_cols = [
            "code", "name", "status", "score",
            "ret20", "ret60", "amount_ma20",
            "amount_ratio_5_20", "max_dd20", "dist_to_60d_high",
        ]

        preview_cols = [c for c in preview_cols if c in selected_df.columns]

        print("\n========== 强趋势池 / 趋势观察池预览 ==========")
        print(selected_df[preview_cols].head(30).to_string(index=False, formatters={
            "ret20": "{:.2%}".format,
            "ret60": "{:.2%}".format,
            "amount_ma20": "{:,.0f}".format,
            "amount_ratio_5_20": "{:.2f}".format,
            "max_dd20": "{:.2%}".format,
            "dist_to_60d_high": "{:.2%}".format,
        }))
    else:
        print("\n指定日期没有股票进入 强趋势池 / 趋势观察池。")
# def analyze_pool_classification(
#     signal_date: pd.Timestamp,
#     stock_data: Dict[str, pd.DataFrame],
#     bench_df: pd.DataFrame,
#     code_name_map: Dict[str, str],
# ) -> None:
#     signal_date = pd.Timestamp(signal_date).normalize()

#     print("\n========== 指定日期股票分类 ==========")
#     print(f"信号日期: {signal_date.date()}")
#     print("功能: 仅分类，不计算未来收益")
#     print("====================================")

#     pool_df = build_pool_on_date(
#         signal_date=signal_date,
#         stock_data=stock_data,
#         bench_df=bench_df,
#         code_name_map=code_name_map,
#     )

#     if pool_df is None or pool_df.empty:
#         print("指定日期未能构建股票分类结果。可能原因：")
#         print("1. 该日期不是交易日；")
#         print("2. 本地数据不足；")
#         print("3. 当天股票没有交易数据；")
#         print("4. 指标窗口不足。")
#         return

#     os.makedirs(OUTPUT_DIR, exist_ok=True)

#     output_xlsx = os.path.join(
#         OUTPUT_DIR,
#         f"{CLASSIFICATION_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.xlsx",
#     )

#     output_csv = os.path.join(
#         OUTPUT_DIR,
#         f"{CLASSIFICATION_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.csv",
#     )

#     try:
#         save_classification_excel(
#             output_xlsx=output_xlsx,
#             signal_date=signal_date,
#             pool_df=pool_df,
#         )

#         print(f"\n分类结果 Excel 已保存到：{output_xlsx}")

#     except Exception as e:
#         print(f"保存分类 Excel 失败：{e}")
#         pool_df.to_csv(output_csv, index=False, encoding="utf-8-sig")
#         print(f"已改为保存 CSV：{output_csv}")

#     summary_df = build_classification_summary(pool_df)

#     if not summary_df.empty:
#         print("\n========== 分类统计 ==========")
#         print(summary_df.to_string(index=False, formatters={
#             "占比": "{:.2%}".format,
#             "平均评分": "{:.2f}".format,
#             "最高评分": "{:.2f}".format,
#             "平均20日涨幅": "{:.2%}".format,
#             "平均60日涨幅": "{:.2%}".format,
#             "平均20日成交额": "{:,.0f}".format,
#             "平均量能比": "{:.2f}".format,
#             "平均20日最大回撤": "{:.2%}".format,
#         }))

#     selected_df = pool_df[pool_df["status"].isin(["强趋势池", "趋势观察池"])].copy()
#     selected_df = selected_df.sort_values("score", ascending=False).reset_index(drop=True)

#     if not selected_df.empty:
#         preview_cols = [
#             "code", "name", "status", "score",
#             "ret20", "ret60", "amount_ma20",
#             "amount_ratio_5_20", "max_dd20", "dist_to_60d_high",
#         ]

#         preview_cols = [c for c in preview_cols if c in selected_df.columns]

#         print("\n========== 强趋势池 / 趋势观察池预览 ==========")
#         print(selected_df[preview_cols].head(30).to_string(index=False, formatters={
#             "ret20": "{:.2%}".format,
#             "ret60": "{:.2%}".format,
#             "amount_ma20": "{:,.0f}".format,
#             "amount_ratio_5_20": "{:.2f}".format,
#             "max_dd20": "{:.2%}".format,
#             "dist_to_60d_high": "{:.2%}".format,
#         }))
#     else:
#         print("\n指定日期没有股票进入 强趋势池 / 趋势观察池。")


# ==============================
# 8.2 回测 Excel 输出
# ==============================

def rename_summary_columns(df: pd.DataFrame, group_col: Optional[str] = None) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()

    out = df.copy()

    rename_map = {
        "horizon": "持有天数",
        "count": "样本数",
        "avg_ret": "平均收益",
        "median_ret": "中位数收益",
        "win_rate": "胜率",
        "bench_ret": "沪深300收益",
        "avg_excess": "平均超额",
        "outperform_rate": "跑赢沪深300比例",
        "status": "分组",
    }

    if group_col and group_col in out.columns:
        rename_map[group_col] = "分组"

    return out.rename(columns=rename_map)


def make_dynamic_exit_readable(dynamic_exit_df: pd.DataFrame, forward_days: int) -> pd.DataFrame:
    if dynamic_exit_df is None or dynamic_exit_df.empty:
        return pd.DataFrame()

    df = dynamic_exit_df.copy()
    key_horizons = get_key_horizons(forward_days)

    keep_cols = [
        "code", "name", "status", "score", "signal_date",
        "dynamic_buy_date", "dynamic_buy_price",
        "dynamic_sell_date", "dynamic_sell_price",
        "dynamic_sell_day", "dynamic_exit_reason",
        "dynamic_final_ret",
        f"dynamic_max_ret_{forward_days}",
        f"dynamic_min_ret_{forward_days}",
        f"dynamic_max_dd_{forward_days}",
        "ret20", "ret60", "amount_ma20",
        "amount_ratio_5_20", "max_dd20", "dist_to_60d_high",
    ]

    for h in key_horizons:
        c = f"strategy_D{h}"
        if c in df.columns:
            keep_cols.append(c)

    keep_cols = [c for c in keep_cols if c in df.columns]
    out = df[keep_cols].copy()

    rename_map = {
        "code": "代码",
        "name": "名称",
        "status": "分组",
        "score": "评分",
        "signal_date": "信号日",
        "dynamic_buy_date": "买入日",
        "dynamic_buy_price": "买入价",
        "dynamic_sell_date": "卖出日",
        "dynamic_sell_price": "卖出价",
        "dynamic_sell_day": "持有天数",
        "dynamic_exit_reason": "卖出原因",
        "dynamic_final_ret": "动态卖出最终收益",
        f"dynamic_max_ret_{forward_days}": f"{forward_days}日内最高浮盈",
        f"dynamic_min_ret_{forward_days}": f"{forward_days}日内最低浮亏",
        f"dynamic_max_dd_{forward_days}": f"{forward_days}日内策略最大回撤",
        "ret20": "信号前20日涨幅",
        "ret60": "信号前60日涨幅",
        "amount_ma20": "20日均成交额",
        "amount_ratio_5_20": "量能比5/20",
        "max_dd20": "信号前20日最大回撤",
        "dist_to_60d_high": "距60日高点",
    }

    for h in key_horizons:
        rename_map[f"strategy_D{h}"] = f"动态D{h}收益"

    out = out.rename(columns=rename_map)

    if "动态卖出最终收益" in out.columns:
        out = out.sort_values("动态卖出最终收益", ascending=False).reset_index(drop=True)

    return out


def make_dynamic_exit_summary_readable(dynamic_exit_summary_df: pd.DataFrame) -> pd.DataFrame:
    if dynamic_exit_summary_df is None or dynamic_exit_summary_df.empty:
        return pd.DataFrame()

    out = dynamic_exit_summary_df.copy()

    rename_map = {
        "group": "分组",
        "count": "样本数",
        "avg_final_ret": "平均最终收益",
        "median_final_ret": "中位数最终收益",
        "win_rate": "胜率",
        "avg_hold_days": "平均持有天数",
        "avg_max_ret": "平均持有期最高浮盈",
        "avg_max_dd": "平均策略最大回撤",
    }

    return out.rename(columns=rename_map)


def build_return_matrix_df(detail_df: pd.DataFrame, forward_days: int) -> pd.DataFrame:
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    cols = ["code", "name", "status", "score", "buy_date", "buy_price"]
    cols = [c for c in cols if c in stock_rows.columns]

    for h in range(1, forward_days + 1):
        c = f"D{h}"
        if c in stock_rows.columns:
            cols.append(c)

    out = stock_rows[cols].copy()

    rename_map = {
        "code": "代码",
        "name": "名称",
        "status": "分组",
        "score": "评分",
        "buy_date": "买入日",
        "buy_price": "买入价",
    }

    return out.rename(columns=rename_map)


def build_excess_matrix_df(excess_df: pd.DataFrame, forward_days: int) -> pd.DataFrame:
    if excess_df is None or excess_df.empty:
        return pd.DataFrame()

    cols = ["code", "name", "status", "score", "buy_date", "avg_excess", "outperform_days", "valid_days"]
    cols = [c for c in cols if c in excess_df.columns]

    for h in range(1, forward_days + 1):
        c = f"excess_D{h}"
        if c in excess_df.columns:
            cols.append(c)

    out = excess_df[cols].copy()

    rename_map = {
        "code": "代码",
        "name": "名称",
        "status": "分组",
        "score": "评分",
        "buy_date": "买入日",
        "avg_excess": "平均超额",
        "outperform_days": "跑赢天数",
        "valid_days": "有效天数",
    }

    for h in range(1, forward_days + 1):
        rename_map[f"excess_D{h}"] = f"D{h}"

    out = out.rename(columns=rename_map)

    if "平均超额" in out.columns:
        out = out.sort_values("平均超额", ascending=False).reset_index(drop=True)

    return out


def build_friendly_selected_df(
    detail_df: pd.DataFrame,
    excess_df: pd.DataFrame,
    forward_days: int,
) -> pd.DataFrame:
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    if stock_rows.empty:
        return pd.DataFrame()

    key_horizons = get_key_horizons(forward_days)
    n = forward_days

    base_cols = [
        "code", "name", "status", "score",
        "signal_date", "buy_date", "buy_price",
        "close_on_signal", "ret20", "ret60",
        "amount_ma20", "amount_ratio_5_20",
        "max_dd20", "dist_to_60d_high",
    ]

    keep_cols = [c for c in base_cols if c in stock_rows.columns]

    for h in key_horizons:
        col = f"D{h}"
        if col in stock_rows.columns:
            keep_cols.append(col)

    stat_cols = [
        f"max_ret_{n}", f"min_ret_{n}", f"max_dd_{n}",
        "day_to_max_ret", "first_profit_5_day",
        "first_profit_10_day", "first_drawdown_5_day",
    ]

    for c in stat_cols:
        if c in stock_rows.columns:
            keep_cols.append(c)

    out = stock_rows[keep_cols].copy()

    if excess_df is not None and not excess_df.empty:
        ex_cols = ["code", "avg_excess", "outperform_days", "valid_days"]

        for h in key_horizons:
            c = f"excess_D{h}"
            if c in excess_df.columns:
                ex_cols.append(c)

        ex_cols = [c for c in ex_cols if c in excess_df.columns]

        out = out.merge(
            excess_df[ex_cols],
            on="code",
            how="left",
        )

    rename_map = {
        "code": "代码",
        "name": "名称",
        "status": "分组",
        "score": "评分",
        "signal_date": "信号日",
        "buy_date": "买入日",
        "buy_price": "买入价",
        "close_on_signal": "信号日收盘",
        "ret20": "信号前20日涨幅",
        "ret60": "信号前60日涨幅",
        "amount_ma20": "20日均成交额",
        "amount_ratio_5_20": "量能比5/20",
        "max_dd20": "信号前20日最大回撤",
        "dist_to_60d_high": "距60日高点",
        f"max_ret_{n}": f"{n}日最高收益",
        f"min_ret_{n}": f"{n}日最低收益",
        f"max_dd_{n}": f"{n}日最大回撤",
        "day_to_max_ret": "最高收益出现在第几天",
        "first_profit_5_day": "首次盈利5%天数",
        "first_profit_10_day": "首次盈利10%天数",
        "first_drawdown_5_day": "首次回撤5%天数",
        "avg_excess": "平均超额收益",
        "outperform_days": "跑赢沪深300天数",
        "valid_days": "有效天数",
    }

    for h in key_horizons:
        rename_map[f"D{h}"] = f"D{h}收益"
        rename_map[f"excess_D{h}"] = f"D{h}超额"

    out = out.rename(columns=rename_map)

    sort_col = f"D{min(30, forward_days)}收益"

    if sort_col in out.columns:
        out = out.sort_values(sort_col, ascending=False).reset_index(drop=True)
    elif "平均超额收益" in out.columns:
        out = out.sort_values("平均超额收益", ascending=False).reset_index(drop=True)

    return out


def save_friendly_excel(
    output_xlsx: str,
    signal_date: pd.Timestamp,
    detail_df: pd.DataFrame,
    excess_df: pd.DataFrame,
    group_summary_df: pd.DataFrame,
    overall_summary_df: pd.DataFrame,
    selected_df: pd.DataFrame,
    pool_df: pd.DataFrame,
    date_map_df: pd.DataFrame,
    forward_days: int,
    dynamic_exit_df: Optional[pd.DataFrame] = None,
    dynamic_exit_summary_df: Optional[pd.DataFrame] = None,
    dynamic_exit_reason_df: Optional[pd.DataFrame] = None,
) -> None:
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    dashboard_rows = []
    dashboard_rows.append(["信号日期", signal_date.strftime("%Y-%m-%d"), "策略信号生成日期"])
    dashboard_rows.append(["买入价格口径", BUY_PRICE_FIELD, "open=买入日开盘价，close=买入日收盘价"])
    dashboard_rows.append(["未来观察天数", forward_days, f"D1~D{forward_days}"])
    dashboard_rows.append(["入选股票数", len(stock_rows), "进入目标池的股票数量"])
    dashboard_rows.append(["动态卖出", "开启" if DYNAMIC_EXIT_ENABLED else "关闭", "股票走弱则卖出，否则持有到期"])
    dashboard_rows.append(["动态卖出最少持有天数", MIN_HOLD_DAYS, "硬止损和跌破MA60不受此限制"])
    dashboard_rows.append(["动态卖出执行时点", EXIT_SELL_TIMING, "close=收盘卖；next_open=次日开盘卖"])
    dashboard_rows.append(["硬止损", STOP_LOSS_PCT, "相对买入价亏损达到该比例卖出"])
    dashboard_rows.append(["移动止盈启动", TRAILING_START_PCT, "盈利达到该比例后启动移动止盈"])
    dashboard_rows.append(["移动止盈回撤", TRAILING_DRAWDOWN_PCT, "从最高收盘价回撤该比例卖出"])

    if overall_summary_df is not None and not overall_summary_df.empty:
        for _, r in overall_summary_df[overall_summary_df["horizon"].isin(get_key_horizons(forward_days))].iterrows():
            h = int(r["horizon"])
            dashboard_rows.append([f"D{h}平均收益", r["avg_ret"], "固定持有口径"])
            dashboard_rows.append([f"D{h}胜率", r["win_rate"], "固定持有口径"])
            dashboard_rows.append([f"D{h}平均超额收益", r["avg_excess"], "固定持有平均收益 - 沪深300"])

    if dynamic_exit_summary_df is not None and not dynamic_exit_summary_df.empty:
        all_row = dynamic_exit_summary_df[dynamic_exit_summary_df["group"] == "全部"]

        if not all_row.empty:
            r = all_row.iloc[0]
            dashboard_rows.append(["动态卖出平均最终收益", r.get("avg_final_ret", np.nan), "动态卖出后的最终收益均值"])
            dashboard_rows.append(["动态卖出胜率", r.get("win_rate", np.nan), "动态最终收益大于0的比例"])
            dashboard_rows.append(["动态卖出平均持有天数", r.get("avg_hold_days", np.nan), "动态卖出平均持有交易日数"])

    dashboard_df = pd.DataFrame(dashboard_rows, columns=["指标", "数值", "说明"])

    friendly_selected_df = build_friendly_selected_df(detail_df, excess_df, forward_days)
    return_matrix_df = build_return_matrix_df(detail_df, forward_days)
    excess_matrix_df = build_excess_matrix_df(excess_df, forward_days)

    readable_overall_summary_df = rename_summary_columns(overall_summary_df)
    readable_group_summary_df = rename_summary_columns(group_summary_df, group_col="status")

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        dashboard_df.to_excel(writer, sheet_name="看板", index=False)
        friendly_selected_df.to_excel(writer, sheet_name="入选股简表", index=False)
        return_matrix_df.to_excel(writer, sheet_name="收益路径", index=False)
        excess_matrix_df.to_excel(writer, sheet_name="超额路径", index=False)
        readable_overall_summary_df.to_excel(writer, sheet_name="整体统计", index=False)
        readable_group_summary_df.to_excel(writer, sheet_name="分组统计", index=False)

        if dynamic_exit_df is not None and not dynamic_exit_df.empty:
            dynamic_exit_readable_df = make_dynamic_exit_readable(dynamic_exit_df, forward_days)
            dynamic_exit_readable_df.to_excel(writer, sheet_name="动态卖出明细", index=False)

        if dynamic_exit_summary_df is not None and not dynamic_exit_summary_df.empty:
            dynamic_exit_summary_readable_df = make_dynamic_exit_summary_readable(dynamic_exit_summary_df)
            dynamic_exit_summary_readable_df.to_excel(writer, sheet_name="动态卖出统计", index=False)

        if dynamic_exit_reason_df is not None and not dynamic_exit_reason_df.empty:
            dynamic_exit_reason_df.to_excel(writer, sheet_name="卖出原因统计", index=False)

        detail_df.to_excel(writer, sheet_name="原始走势明细", index=False)
        excess_df.to_excel(writer, sheet_name="原始超额明细", index=False)
        selected_df.to_excel(writer, sheet_name="入选股票原始", index=False)
        pool_df.to_excel(writer, sheet_name="全市场分类", index=False)
        date_map_df.to_excel(writer, sheet_name="日映射", index=False)

        wb = writer.book

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            if sheet_name == "看板":
                style_worksheet(ws, freeze="A2", apply_red_green=True)
                ws.column_dimensions["A"].width = 26
                ws.column_dimensions["B"].width = 20
                ws.column_dimensions["C"].width = 56
            elif sheet_name in ["收益路径", "超额路径"]:
                style_worksheet(ws, freeze="G2", apply_red_green=True)
            else:
                style_worksheet(ws, freeze="A2", apply_red_green=True)

        for raw_sheet in ["原始走势明细", "原始超额明细"]:
            if raw_sheet in wb.sheetnames:
                wb[raw_sheet].sheet_state = "hidden"


# ==============================
# 9. 主分析函数：未来回测
# ==============================

def build_dynamic_exit_detail(
    selected_df: pd.DataFrame,
    stock_data: Dict[str, pd.DataFrame],
    forward_dates: List[pd.Timestamp],
    forward_days: int,
) -> pd.DataFrame:
    if selected_df is None or selected_df.empty:
        return pd.DataFrame()

    rows = []

    for _, stock in selected_df.iterrows():
        code = str(stock["code"])

        if code not in stock_data:
            continue

        path = calc_forward_trade_with_dynamic_exit(
            df=stock_data[code],
            forward_dates=forward_dates,
            buy_price_field=BUY_PRICE_FIELD,
            min_hold_days=MIN_HOLD_DAYS,
            exit_sell_timing=EXIT_SELL_TIMING,
            stop_loss_pct=STOP_LOSS_PCT,
            trailing_start_pct=TRAILING_START_PCT,
            trailing_drawdown_pct=TRAILING_DRAWDOWN_PCT,
        )

        if not path:
            continue

        row = {
            "code": code,
            "name": stock.get("name", ""),
            "status": stock.get("status", ""),
            "score": stock.get("score", np.nan),
            "signal_date": stock.get("signal_date", ""),
            "close_on_signal": stock.get("close", np.nan),
            "ret20": stock.get("ret20", np.nan),
            "ret60": stock.get("ret60", np.nan),
            "amount_ma20": stock.get("amount_ma20", np.nan),
            "amount_ratio_5_20": stock.get("amount_ratio_5_20", np.nan),
            "max_dd20": stock.get("max_dd20", np.nan),
            "dist_to_60d_high": stock.get("dist_to_60d_high", np.nan),
        }

        row.update(path)
        rows.append(row)

    if not rows:
        return pd.DataFrame()

    df = pd.DataFrame(rows)

    if "dynamic_final_ret" in df.columns:
        df = df.sort_values(["dynamic_final_ret", "score"], ascending=[False, False])

    return df.reset_index(drop=True)


def analyze_pool_forward_returns(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Dict[str, str],
    forward_days: int,
    statuses: List[str],
) -> None:
    signal_date = pd.Timestamp(signal_date).normalize()
    print("\n========== 指定日期股票池未来走势分析 ==========")
    print(f"信号日期: {signal_date.date()}")
    print(f"分析池: {statuses}")
    print(f"未来交易日数: {forward_days}")
    print("==============================================")

    pool_df = build_pool_on_date(
        signal_date=signal_date,
        stock_data=stock_data,
        bench_df=bench_df,
        code_name_map=code_name_map,
    )

    if pool_df.empty:
        print("指定日期未能构建股票池。")
        return

    selected_df = pool_df[pool_df["status"].isin(statuses)].copy()
    selected_df = selected_df.sort_values("score", ascending=False).reset_index(drop=True)

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    if selected_df.empty:
        print(f"{signal_date.date()} 没有股票进入 {statuses}。")

        output_all_pool = os.path.join(
            OUTPUT_DIR,
            f"{FORWARD_ANALYSIS_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}_all_pool.csv",
        )

        pool_df.to_csv(output_all_pool, index=False, encoding="utf-8-sig")
        print(f"已保存当日全市场分类到：{output_all_pool}")

        return

    print(f"入选股票数量: {len(selected_df)}")

    preview_cols = [
        "code", "name", "status", "score", "ret20", "ret60",
        "amount_ma20", "amount_ratio_5_20",
    ]

    preview_cols = [c for c in preview_cols if c in selected_df.columns]

    print("\n入选股票预览：")
    print(selected_df[preview_cols].head(20).to_string(index=False))

    forward_dates = get_forward_trading_dates(
        signal_date=signal_date,
        bench_df=bench_df,
        forward_days=forward_days,
    )

    if len(forward_dates) < forward_days:
        print(f"基准未来交易日不足：需要 {forward_days} 天，实际只有 {len(forward_dates)} 天。")
        print("请降低 --days，或者选择更早的信号日期。")
        return

    date_map_df = pd.DataFrame({
        "horizon": list(range(1, forward_days + 1)),
        "date": [d.strftime("%Y-%m-%d") for d in forward_dates],
    })

    bench_path = calc_forward_return_path_by_calendar(
        df=bench_df,
        forward_dates=forward_dates,
        buy_price_field=BUY_PRICE_FIELD,
    )

    if not bench_path:
        print("基准未来走势计算失败。")
        return

    detail_rows = []

    bench_row = {
        "row_type": "benchmark",
        "code": "沪深300",
        "name": "沪深300",
        "status": "benchmark",
        "score": np.nan,
        "signal_date": signal_date.strftime("%Y-%m-%d"),
    }

    bench_row.update(bench_path)
    detail_rows.append(bench_row)

    skipped = 0

    for _, stock in selected_df.iterrows():
        code = str(stock["code"])

        if code not in stock_data:
            skipped += 1
            continue

        path = calc_forward_return_path_by_calendar(
            df=stock_data[code],
            forward_dates=forward_dates,
            buy_price_field=BUY_PRICE_FIELD,
        )

        if not path:
            skipped += 1
            continue

        row = {
            "row_type": "stock",
            "code": code,
            "name": stock.get("name", code_name_map.get(code, "")),
            "status": stock["status"],
            "score": stock["score"],
            "signal_date": signal_date.strftime("%Y-%m-%d"),

            "close_on_signal": stock["close"],
            "ret20": stock["ret20"],
            "ret60": stock["ret60"],
            "ret20_rank_pct": stock["ret20_rank_pct"],
            "ret60_rank_pct": stock["ret60_rank_pct"],
            "amount_ma20": stock["amount_ma20"],
            "amount_ratio_5_20": stock["amount_ratio_5_20"],
            "max_dd20": stock["max_dd20"],
            "dist_to_60d_high": stock["dist_to_60d_high"],
        }

        row.update(path)
        detail_rows.append(row)

    detail_df = pd.DataFrame(detail_rows)

    if len(detail_df) <= 1:
        print("入选股票未来数据不足，无法生成走势分析。")
        return

    if skipped > 0:
        print(f"有 {skipped} 只股票由于买入日无交易、价格异常或未来数据缺失被跳过。")

    excess_df = build_excess_detail(
        detail_df=detail_df,
        forward_days=forward_days,
    )

    group_summary_df = build_forward_summary(
        detail_df=detail_df,
        forward_days=forward_days,
        group_col="status",
    )

    overall_summary_df = build_overall_forward_summary(
        detail_df=detail_df,
        forward_days=forward_days,
    )

    dynamic_exit_df = pd.DataFrame()
    dynamic_exit_summary_df = pd.DataFrame()
    dynamic_exit_reason_df = pd.DataFrame()

    if DYNAMIC_EXIT_ENABLED:
        dynamic_exit_df = build_dynamic_exit_detail(
            selected_df=selected_df,
            stock_data=stock_data,
            forward_dates=forward_dates,
            forward_days=forward_days,
        )

        dynamic_exit_summary_df = build_dynamic_exit_summary(
            dynamic_exit_df=dynamic_exit_df,
            group_col="status",
        )

        dynamic_exit_reason_df = build_dynamic_exit_reason_summary(
            dynamic_exit_df=dynamic_exit_df,
        )

    output_xlsx = os.path.join(
        OUTPUT_DIR,
        f"{FORWARD_ANALYSIS_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.xlsx",
    )

    output_detail_csv = os.path.join(
        OUTPUT_DIR,
        f"{FORWARD_ANALYSIS_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}_detail.csv",
    )

    try:
        save_friendly_excel(
            output_xlsx=output_xlsx,
            signal_date=signal_date,
            detail_df=detail_df,
            excess_df=excess_df,
            group_summary_df=group_summary_df,
            overall_summary_df=overall_summary_df,
            selected_df=selected_df,
            pool_df=pool_df,
            date_map_df=date_map_df,
            forward_days=forward_days,
            dynamic_exit_df=dynamic_exit_df,
            dynamic_exit_summary_df=dynamic_exit_summary_df,
            dynamic_exit_reason_df=dynamic_exit_reason_df,
        )

        print(f"\n友好版分析结果已保存到：{output_xlsx}")

    except Exception as e:
        print(f"保存友好版 Excel 失败：{e}")
        print("改为保存 CSV 明细。")

        detail_df.to_csv(output_detail_csv, index=False, encoding="utf-8-sig")
        print(f"走势明细已保存到：{output_detail_csv}")

    print("\n========== 未来走势摘要 ==========")

    key_horizons = [1, 3, 5, 10, 20, 30]
    key_horizons = [h for h in key_horizons if h <= forward_days]

    if not overall_summary_df.empty:
        display_df = overall_summary_df[
            overall_summary_df["horizon"].isin(key_horizons)
        ].copy()

        if not display_df.empty:
            print("\n固定持有整体统计：")
            print(display_df.to_string(index=False, formatters={
                "avg_ret": "{:.2%}".format,
                "median_ret": "{:.2%}".format,
                "win_rate": "{:.2%}".format,
                "bench_ret": "{:.2%}".format,
                "avg_excess": "{:.2%}".format,
                "outperform_rate": "{:.2%}".format,
            }))

    if not group_summary_df.empty:
        display_group_df = group_summary_df[
            group_summary_df["horizon"].isin(key_horizons)
        ].copy()

        if not display_group_df.empty:
            print("\n固定持有分组统计：")
            print(display_group_df.to_string(index=False, formatters={
                "avg_ret": "{:.2%}".format,
                "median_ret": "{:.2%}".format,
                "win_rate": "{:.2%}".format,
                "bench_ret": "{:.2%}".format,
                "avg_excess": "{:.2%}".format,
                "outperform_rate": "{:.2%}".format,
            }))

    if DYNAMIC_EXIT_ENABLED and dynamic_exit_summary_df is not None and not dynamic_exit_summary_df.empty:
        print("\n动态卖出统计：")
        print(dynamic_exit_summary_df.to_string(index=False, formatters={
            "avg_final_ret": "{:.2%}".format,
            "median_final_ret": "{:.2%}".format,
            "win_rate": "{:.2%}".format,
            "avg_hold_days": "{:.2f}".format,
            "avg_max_ret": "{:.2%}".format,
            "avg_max_dd": "{:.2%}".format,
        }))


# ==============================
# 10. 命令行入口
# ==============================

def parse_args():
    parser = argparse.ArgumentParser(
        description="指定日期股票分类 / 股票池未来走势分析"
    )

    parser.add_argument(
        "--date",
        required=True,
        help="信号日期，例如 2026-04-30。",
    )

    parser.add_argument(
        "--days",
        type=int,
        default=30,
        help="统计未来多少个交易日，默认 30。分类模式下无效。",
    )

    parser.add_argument(
        "--classify-only",
        action="store_true",
        help="只对指定日期股票进行分类，不做未来走势回测。",
    )

    parser.add_argument(
        "--statuses",
        type=str,
        default="强趋势池,趋势观察池",
        help="要分析的股票池，逗号分隔，默认：强趋势池,趋势观察池。",
    )

    parser.add_argument(
        "--use-prev-trading-day",
        action="store_true",
        help="如果输入日期不是交易日，则自动使用前一个交易日。",
    )

    parser.add_argument(
        "--cache-dir",
        type=str,
        default="data",
        help="本地数据目录，默认 data。",
    )

    parser.add_argument(
        "--output-dir",
        type=str,
        default="output",
        help="输出目录，默认 output。",
    )

    parser.add_argument(
        "--universe-source",
        type=str,
        default="local",
        choices=["local", "ak", "auto"],
        help=(
            "股票池来源："
            "local=纯本地扫描 data/*.csv，不联网；"
            "ak=联网用 AkShare 获取代码名称并和本地 CSV 求交集；"
            "auto=优先 AkShare，失败后回退本地。默认 local。"
        ),
    )

    parser.add_argument(
        "--min-bars",
        type=int,
        default=130,
        help="最少K线数量，默认 130。",
    )

    parser.add_argument(
        "--min-score",
        type=int,
        default=75,
        help="趋势观察池最低评分，默认 75。",
    )

    parser.add_argument(
        "--min-amount",
        type=float,
        default=80_000_000,
        help="20日平均成交额门槛，默认 80000000。",
    )

    parser.add_argument(
        "--buy-price-field",
        type=str,
        default="open",
        choices=["open", "close"],
        help="买入基准价格字段，默认 open。",
    )

    parser.add_argument(
        "--no-dynamic-exit",
        action="store_true",
        help="关闭动态卖出回测。默认开启。",
    )

    parser.add_argument(
        "--min-hold-days",
        type=int,
        default=3,
        help="动态卖出最少持有天数，默认 3。",
    )

    parser.add_argument(
        "--exit-sell-timing",
        type=str,
        default="next_open",
        choices=["close", "next_open"],
        help="动态卖出执行时点：close=触发日收盘卖出；next_open=下一交易日开盘卖出。默认 next_open。",
    )

    parser.add_argument(
        "--stop-loss",
        type=float,
        default=-0.08,
        help="动态卖出硬止损比例，默认 -0.08，即 -8%。",
    )

    parser.add_argument(
        "--trailing-start",
        type=float,
        default=0.10,
        help="移动止盈启动收益，默认 0.10，即盈利10%后启动。",
    )

    parser.add_argument(
        "--trailing-drawdown",
        type=float,
        default=0.08,
        help="移动止盈允许从高点回撤比例，默认 0.08，即从高点回撤8%卖出。",
    )

    return parser.parse_args()


def main():
    global CACHE_DIR
    global OUTPUT_DIR
    global BENCHMARK_CACHE_FILE
    global STOCK_NAME_FILE
    global MIN_BARS
    global MIN_SCORE
    global MIN_AVG_AMOUNT_20
    global BUY_PRICE_FIELD

    global DYNAMIC_EXIT_ENABLED
    global MIN_HOLD_DAYS
    global EXIT_SELL_TIMING
    global STOP_LOSS_PCT
    global TRAILING_START_PCT
    global TRAILING_DRAWDOWN_PCT

    args = parse_args()

    CACHE_DIR = args.cache_dir
    OUTPUT_DIR = args.output_dir
    BENCHMARK_CACHE_FILE = os.path.join(CACHE_DIR, "_benchmark_000300.csv")
    STOCK_NAME_FILE = os.path.join(CACHE_DIR, "_stock_names.csv")

    MIN_BARS = args.min_bars
    MIN_SCORE = args.min_score
    MIN_AVG_AMOUNT_20 = args.min_amount
    BUY_PRICE_FIELD = args.buy_price_field

    DYNAMIC_EXIT_ENABLED = not args.no_dynamic_exit
    MIN_HOLD_DAYS = args.min_hold_days
    EXIT_SELL_TIMING = args.exit_sell_timing
    STOP_LOSS_PCT = args.stop_loss
    TRAILING_START_PCT = args.trailing_start
    TRAILING_DRAWDOWN_PCT = args.trailing_drawdown

    statuses = args.statuses.replace("，", ",")
    statuses = [s.strip() for s in statuses.split(",") if s.strip()]

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    print("========== 指定日期股票分类 / 未来走势分析 ==========")
    print(f"输入信号日期: {args.date}")
    print(f"运行模式: {'仅分类' if args.classify_only else '回测分析'}")
    print(f"未来交易日数: {args.days}")
    print(f"分析池: {statuses}")
    print(f"数据目录: {CACHE_DIR}")
    print(f"输出目录: {OUTPUT_DIR}")
    print(f"基准文件: {BENCHMARK_CACHE_FILE}")
    print(f"股票池来源: {args.universe_source}")
    print(f"最少K线数量: {MIN_BARS}")
    print(f"最低评分: {MIN_SCORE}")
    print(f"20日平均成交额门槛: {MIN_AVG_AMOUNT_20:,.0f}")
    print(f"买入价格字段: {BUY_PRICE_FIELD}")

    print(f"动态卖出: {'开启' if DYNAMIC_EXIT_ENABLED else '关闭'}")
    print(f"动态卖出最少持有天数: {MIN_HOLD_DAYS}")
    print(f"动态卖出执行时点: {EXIT_SELL_TIMING}")
    print(f"硬止损: {STOP_LOSS_PCT:.2%}")
    print(f"移动止盈启动: {TRAILING_START_PCT:.2%}")
    print(f"移动止盈回撤: {TRAILING_DRAWDOWN_PCT:.2%}")

    if args.universe_source == "local":
        print("注意：当前为纯本地模式，不联网。")
    elif args.universe_source == "ak":
        print("注意：当前会联网获取股票代码和名称，但不会下载历史行情。")
    elif args.universe_source == "auto":
        print("注意：当前会尝试联网获取股票代码和名称，失败后回退本地；不会下载历史行情。")

    print("====================================================")

    print("\n构建股票池...")
    try:
        universe = get_universe(CACHE_DIR, source=args.universe_source)
    except Exception as e:
        print(f"构建股票池失败：{e}")
        return

    if universe is None or universe.empty:
        print("股票池为空，程序终止。")
        return

    codes = universe["code"].astype(str).apply(normalize_code).tolist()
    code_name_map = build_code_name_map(universe)

    print(f"最终可分析本地股票数量: {len(codes)}")
    print(f"有名称的股票数量: {sum(1 for c in codes if code_name_map.get(c, ''))}")

    print("\n加载本地基准数据...")
    try:
        bench_df = fetch_benchmark_data()
    except Exception as e:
        print(f"加载基准数据失败：{e}")
        return

    try:
        signal_date = resolve_signal_date(
            requested_date=args.date,
            bench_df=bench_df,
            use_prev_trading_day=args.use_prev_trading_day,
        )
    except Exception as e:
        print(f"信号日期错误：{e}")
        return

    print("\n加载本地股票数据...")
    stock_data = preload_all_stock_data(codes)

    if not stock_data:
        print("没有成功加载任何股票数据，程序终止。")
        return

    if args.classify_only:
        analyze_pool_classification(
            signal_date=signal_date,
            stock_data=stock_data,
            bench_df=bench_df,
            code_name_map=code_name_map,
        )
    else:
        analyze_pool_forward_returns(
            signal_date=signal_date,
            stock_data=stock_data,
            bench_df=bench_df,
            code_name_map=code_name_map,
            forward_days=args.days,
            statuses=statuses,
        )


if __name__ == "__main__":
    main()
