#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
A股趋势策略增强版：
指定日期股票分类 + 连续多日评分趋势 + 短线强势池 + 短中线转强池 + 固定持有回测

使用示例：

仅分类：
python pool04_enhanced_score_trend.py --date 2026-06-18 --classify-only

非交易日自动前移：
python pool04_enhanced_score_trend.py --date 2026-06-19 --use-prev-trading-day --classify-only

回测：
python pool04_enhanced_score_trend.py --date 2026-06-18 --days 30

调整评分趋势回看天数：
python pool04_enhanced_score_trend.py --date 2026-06-18 --classify-only --score-trend-lookback 10
"""

import os
import re
import argparse
import warnings
from typing import Dict, List, Optional

import numpy as np
import pandas as pd
from tqdm import tqdm
from numpy.lib.stride_tricks import sliding_window_view

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


# =========================================================
# 0. 默认参数
# =========================================================

CACHE_DIR = "data"
OUTPUT_DIR = "output"

BENCHMARK_CACHE_FILE = os.path.join(CACHE_DIR, "_benchmark_000300.csv")
STOCK_NAME_FILE = os.path.join(CACHE_DIR, "_stock_names.csv")
STOCK_INDUSTRY_FILE = os.path.join(CACHE_DIR, "_stock_industry.csv")

MIN_SCORE = 75
STRONG_SCORE = 85
MIN_BARS = 130
MIN_AVG_AMOUNT_20 = 80_000_000

BUY_PRICE_FIELD = "open"

CLASSIFICATION_OUTPUT_PREFIX = "pool_classification_trend"
FORWARD_ANALYSIS_OUTPUT_PREFIX = "pool_forward"

MARKET_FILTER_ENABLED = True
INDUSTRY_FILTER_ENABLED = True
LIMIT_CHECK_ENABLED = True

BUY_COMMISSION_RATE = 0.0003
SELL_COMMISSION_RATE = 0.0003
STAMP_TAX_RATE = 0.001
SLIPPAGE_RATE = 0.0005

LIMIT_UP_THRESHOLD = 0.095
LIMIT_DOWN_THRESHOLD = -0.095

ATR_PERIOD = 20

SCORE_TREND_LOOKBACK = 10
SCORE_TREND_WINDOW = 5

POOL_STATUSES = [
    "短中线转强池",
    "短线强势池",
    "强趋势池",
    "趋势观察池",
]


# =========================================================
# 1. 工具函数
# =========================================================

def normalize_code(code) -> str:
    s = str(code).strip()
    m = re.search(r"(\d{6})", s)
    if m:
        return m.group(1)
    return s.zfill(6)


def is_mainboard_code(code: str) -> bool:
    code = normalize_code(code)
    return len(code) == 6 and code.isdigit() and code.startswith(("60", "00"))


def is_bad_name(name: str) -> bool:
    if pd.isna(name):
        return False
    s = str(name).strip()
    if not s:
        return False
    return bool(re.search(r"ST|\*ST|退", s, flags=re.IGNORECASE))


def normalize_date_col(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["date"] = pd.to_datetime(df["date"], errors="coerce").dt.normalize()
    df = df.dropna(subset=["date"])
    df = df.sort_values("date").drop_duplicates("date", keep="last").reset_index(drop=True)
    return df


def max_drawdown(series: pd.Series) -> float:
    s = series.dropna()
    if s.empty:
        return np.nan
    dd = s / s.cummax() - 1
    return float(dd.min())


def rolling_max_drawdown_np(close: pd.Series, window: int = 20) -> np.ndarray:
    arr = close.to_numpy(dtype="float64", copy=False)
    n = len(arr)
    out = np.full(n, np.nan)

    if n < window:
        return out

    windows = sliding_window_view(arr, window_shape=window)
    valid = ~np.isnan(windows).any(axis=1)
    vals = np.full(windows.shape[0], np.nan)

    if valid.any():
        w = windows[valid]
        cum_max = np.maximum.accumulate(w, axis=1)
        dd = w / cum_max - 1
        vals[valid] = dd.min(axis=1)

    out[window - 1:] = vals
    return out


def safe_bool(x) -> bool:
    if pd.isna(x):
        return False
    return bool(x)


def standardize_local_df(df: pd.DataFrame, code: str = "") -> Optional[pd.DataFrame]:
    if df is None or df.empty:
        return None

    df = df.copy()

    rename_map = {
        "日期": "date",
        "时间": "date",
        "开盘": "open",
        "收盘": "close",
        "最高": "high",
        "最低": "low",
        "成交量": "volume",
        "成交额": "amount",
        "成交金额": "amount",
        "涨跌幅": "pct_chg",
        "涨跌额": "change",
        "换手率": "turnover",
        "振幅": "amplitude",
        "date": "date",
        "open": "open",
        "close": "close",
        "high": "high",
        "low": "low",
        "volume": "volume",
        "amount": "amount",
        "pct_chg": "pct_chg",
        "turnover": "turnover",
        "change": "change",
        "amplitude": "amplitude",
    }

    df = df.rename(columns=rename_map)

    if "date" not in df.columns:
        print(f"{code} 缺少 date 字段，跳过。")
        return None

    required_cols = ["date", "open", "high", "low", "close"]
    missing = [c for c in required_cols if c not in df.columns]

    if missing:
        print(f"{code} 缺少必要字段 {missing}，跳过。")
        return None

    df = normalize_date_col(df)

    numeric_cols = [
        "open", "high", "low", "close", "volume", "amount",
        "pct_chg", "turnover", "change", "amplitude",
    ]

    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df = df.dropna(subset=["open", "high", "low", "close"])

    if df.empty:
        return None

    if "volume" not in df.columns:
        df["volume"] = np.nan

    if "amount" not in df.columns or df["amount"].isna().all():
        if "volume" in df.columns and not df["volume"].isna().all():
            df["amount"] = df["volume"] * df["close"]
        else:
            df["amount"] = np.nan

    if "pct_chg" not in df.columns or df["pct_chg"].isna().all():
        df["pct_chg"] = df["close"].pct_change() * 100

    if "turnover" not in df.columns:
        df["turnover"] = np.nan

    if "change" not in df.columns:
        df["change"] = df["close"].diff()

    return df.reset_index(drop=True)


def net_buy_price(raw_price: float) -> float:
    if pd.isna(raw_price) or raw_price <= 0:
        return np.nan
    return raw_price * (1 + SLIPPAGE_RATE) * (1 + BUY_COMMISSION_RATE)


def net_sell_price(raw_price: float) -> float:
    if pd.isna(raw_price) or raw_price <= 0:
        return np.nan
    return raw_price * (1 - SLIPPAGE_RATE) * (1 - SELL_COMMISSION_RATE - STAMP_TAX_RATE)


def calc_net_return(raw_sell_price: float, raw_buy_price: float) -> float:
    bp = net_buy_price(raw_buy_price)
    sp = net_sell_price(raw_sell_price)

    if pd.isna(bp) or bp <= 0 or pd.isna(sp) or sp <= 0:
        return np.nan

    return sp / bp - 1


def get_row_price(row: pd.Series, field: str, fallback: str = "close") -> float:
    if field in row.index and not pd.isna(row[field]):
        return float(row[field])
    if fallback in row.index and not pd.isna(row[fallback]):
        return float(row[fallback])
    return np.nan


def is_open_limit_up(row: pd.Series) -> bool:
    if not LIMIT_CHECK_ENABLED:
        return False

    pre_close = row.get("pre_close", np.nan)
    open_price = row.get("open", np.nan)

    if pd.isna(pre_close) or pre_close <= 0 or pd.isna(open_price) or open_price <= 0:
        return False

    return open_price / pre_close - 1 >= LIMIT_UP_THRESHOLD


def is_open_limit_down(row: pd.Series) -> bool:
    if not LIMIT_CHECK_ENABLED:
        return False

    pre_close = row.get("pre_close", np.nan)
    open_price = row.get("open", np.nan)

    if pd.isna(pre_close) or pre_close <= 0 or pd.isna(open_price) or open_price <= 0:
        return False

    return open_price / pre_close - 1 <= LIMIT_DOWN_THRESHOLD


# =========================================================
# 2. 股票列表、名称、行业
# =========================================================

def scan_local_stock_codes(data_dir: str = "data") -> List[str]:
    if not os.path.exists(data_dir):
        raise FileNotFoundError(f"数据目录不存在：{data_dir}")

    codes = []

    for fname in os.listdir(data_dir):
        if not fname.lower().endswith(".csv"):
            continue

        stem = fname[:-4]

        if stem.startswith("_benchmark") or stem.startswith("_stock_names") or stem.startswith("_stock_industry"):
            continue

        code = normalize_code(stem)

        if is_mainboard_code(code):
            codes.append(code)

    return sorted(set(codes))


def load_local_stock_names(data_dir: str = "data") -> Dict[str, str]:
    path = os.path.join(data_dir, "_stock_names.csv")

    if not os.path.exists(path):
        return {}

    try:
        df = pd.read_csv(path, dtype=str)
    except Exception as e:
        print(f"读取股票名称文件失败：{e}")
        return {}

    if df is None or df.empty:
        return {}

    code_col = None
    name_col = None

    for c in ["code", "代码", "股票代码", "symbol", "证券代码"]:
        if c in df.columns:
            code_col = c
            break

    for c in ["name", "名称", "股票简称", "stock_name", "证券简称"]:
        if c in df.columns:
            name_col = c
            break

    if code_col is None or name_col is None:
        return {}

    tmp = df[[code_col, name_col]].copy()
    tmp.columns = ["code", "name"]
    tmp["code"] = tmp["code"].apply(normalize_code)
    tmp["name"] = tmp["name"].astype(str).str.strip()
    tmp = tmp[tmp["code"].apply(is_mainboard_code)]
    tmp = tmp[~tmp["name"].apply(is_bad_name)]
    tmp = tmp.drop_duplicates(subset=["code"], keep="last")

    return dict(zip(tmp["code"], tmp["name"]))


def load_local_stock_industry(data_dir: str = "data") -> Dict[str, str]:
    path = os.path.join(data_dir, "_stock_industry.csv")

    if not os.path.exists(path):
        return {}

    try:
        df = pd.read_csv(path, dtype=str)
    except Exception as e:
        print(f"读取行业文件失败：{e}")
        return {}

    if df is None or df.empty:
        return {}

    code_col = None
    industry_col = None

    for c in ["code", "代码", "股票代码", "symbol", "证券代码"]:
        if c in df.columns:
            code_col = c
            break

    for c in ["industry", "行业", "申万行业", "板块", "sector"]:
        if c in df.columns:
            industry_col = c
            break

    if code_col is None or industry_col is None:
        print(f"行业文件字段异常：{path}，需要 code/industry 或 代码/行业")
        return {}

    tmp = df[[code_col, industry_col]].copy()
    tmp.columns = ["code", "industry"]
    tmp["code"] = tmp["code"].apply(normalize_code)
    tmp["industry"] = tmp["industry"].astype(str).str.strip()
    tmp = tmp[tmp["code"].apply(is_mainboard_code)]
    tmp = tmp.drop_duplicates(subset=["code"], keep="last")

    return dict(zip(tmp["code"], tmp["industry"]))


def get_a_stock_universe_from_ak() -> pd.DataFrame:
    print("正在通过 AkShare 获取 A 股股票列表和名称...")

    try:
        import akshare as ak
    except Exception as e:
        raise RuntimeError(
            "未安装 AkShare。如需使用 --universe-source ak，请先执行：pip install akshare"
        ) from e

    stock_info = None
    last_err = None

    try:
        stock_info = ak.stock_info_a_code_name()
    except Exception as e:
        last_err = e
        stock_info = None

    if stock_info is None or stock_info.empty:
        try:
            stock_info = ak.stock_zh_a_spot()
        except Exception as e:
            last_err = e
            stock_info = None

    if stock_info is None or stock_info.empty:
        raise RuntimeError(f"无法通过 AkShare 获取股票列表。最后错误：{last_err}")

    if "code" in stock_info.columns and "name" in stock_info.columns:
        df = stock_info[["code", "name"]].copy()
    elif "代码" in stock_info.columns and "名称" in stock_info.columns:
        df = stock_info[["代码", "名称"]].copy()
        df.columns = ["code", "name"]
    elif "symbol" in stock_info.columns and "name" in stock_info.columns:
        df = stock_info[["symbol", "name"]].copy()
        df.columns = ["code", "name"]
    elif "代码" in stock_info.columns and "简称" in stock_info.columns:
        df = stock_info[["代码", "简称"]].copy()
        df.columns = ["code", "name"]
    else:
        raise ValueError(f"股票列表字段异常：{stock_info.columns.tolist()}")

    df["code"] = df["code"].apply(normalize_code)
    df["name"] = df["name"].astype(str).str.strip()
    df = df[df["code"].apply(is_mainboard_code)]
    df = df[~df["name"].apply(is_bad_name)]
    df = df.drop_duplicates(subset=["code"]).sort_values("code").reset_index(drop=True)

    return df


def save_stock_names(data_dir: str, df: pd.DataFrame) -> None:
    if df is None or df.empty:
        return

    os.makedirs(data_dir, exist_ok=True)
    path = os.path.join(data_dir, "_stock_names.csv")

    out = df[["code", "name"]].copy()
    out["code"] = out["code"].astype(str).apply(normalize_code)
    out["name"] = out["name"].astype(str).str.strip()
    out = out[out["code"].apply(is_mainboard_code)]
    out = out[~out["name"].apply(is_bad_name)]
    out = out.drop_duplicates(subset=["code"], keep="last")
    out = out.sort_values("code").reset_index(drop=True)

    out.to_csv(path, index=False, encoding="utf-8-sig")
    print(f"股票名称已保存到：{path}")


def get_mainboard_stocks_from_local(data_dir: str = "data") -> pd.DataFrame:
    codes = scan_local_stock_codes(data_dir)

    if not codes:
        raise RuntimeError(f"{data_dir} 目录下没有找到主板股票 CSV 文件。")

    name_map = load_local_stock_names(data_dir)

    df = pd.DataFrame({
        "code": codes,
        "name": [name_map.get(c, "") for c in codes],
    })

    if name_map:
        before = len(df)
        df = df[~df["name"].apply(is_bad_name)].copy()
        after = len(df)
        if before != after:
            print(f"根据本地名称文件剔除 ST/退市股票：{before - after} 只")

    df = df.drop_duplicates(subset=["code"]).reset_index(drop=True)
    print(f"从本地 {data_dir} 目录识别主板股票数量：{len(df)}")
    return df


def get_mainboard_stocks_from_ak_and_local(data_dir: str = "data") -> pd.DataFrame:
    local_codes = set(scan_local_stock_codes(data_dir))

    if not local_codes:
        raise RuntimeError(f"{data_dir} 目录下没有找到主板股票 CSV 文件。")

    ak_df = get_a_stock_universe_from_ak()
    save_stock_names(data_dir, ak_df)

    df = ak_df[ak_df["code"].isin(local_codes)].copy()
    df = df.sort_values("code").reset_index(drop=True)

    print(f"本地 CSV 股票数量：{len(local_codes)}")
    print(f"AkShare 主板股票数量：{len(ak_df)}")
    print(f"两者交集数量：{len(df)}")

    if df.empty:
        raise RuntimeError("AkShare 股票列表和本地 CSV 没有交集，请检查代码格式。")

    return df


def get_universe(data_dir: str = "data", source: str = "local") -> pd.DataFrame:
    source = str(source).lower().strip()

    if source == "local":
        return get_mainboard_stocks_from_local(data_dir)

    if source == "ak":
        return get_mainboard_stocks_from_ak_and_local(data_dir)

    if source == "auto":
        try:
            return get_mainboard_stocks_from_ak_and_local(data_dir)
        except Exception as e:
            print(f"AkShare 获取失败，回退本地模式。错误：{e}")
            return get_mainboard_stocks_from_local(data_dir)

    raise ValueError(f"未知 universe-source：{source}")


def build_code_name_map(universe: pd.DataFrame) -> Dict[str, str]:
    if universe is None or universe.empty:
        return {}

    tmp = universe.copy()
    if "code" not in tmp.columns:
        return {}

    name_col = None
    for c in ["name", "名称", "stock_name", "股票简称"]:
        if c in tmp.columns:
            name_col = c
            break

    if name_col is None:
        return {}

    tmp = tmp[["code", name_col]].copy()
    tmp.columns = ["code", "name"]
    tmp["code"] = tmp["code"].astype(str).apply(normalize_code)
    tmp["name"] = tmp["name"].astype(str).str.strip()
    tmp = tmp[tmp["name"] != ""]
    tmp = tmp[tmp["name"].str.lower() != "nan"]

    return dict(zip(tmp["code"], tmp["name"]))


# =========================================================
# 3. 指标计算
# =========================================================

def calc_atr(df: pd.DataFrame, period: int = 20) -> pd.Series:
    high = df["high"]
    low = df["low"]
    close = df["close"]
    pre_close = close.shift(1)

    tr1 = high - low
    tr2 = (high - pre_close).abs()
    tr3 = (low - pre_close).abs()

    tr = pd.concat([tr1, tr2, tr3], axis=1).max(axis=1)
    atr = tr.rolling(period).mean()

    return atr


def add_indicators(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    df["pre_close"] = df["close"].shift(1)

    df["ma5"] = df["close"].rolling(5).mean()
    df["ma10"] = df["close"].rolling(10).mean()
    df["ma20"] = df["close"].rolling(20).mean()
    df["ma60"] = df["close"].rolling(60).mean()
    df["ma120"] = df["close"].rolling(120).mean()

    df["ma5_slope_3"] = df["ma5"] / df["ma5"].shift(3) - 1
    df["ma10_slope_5"] = df["ma10"] / df["ma10"].shift(5) - 1
    df["ma20_slope_5"] = df["ma20"] / df["ma20"].shift(5) - 1
    df["ma60_slope_10"] = df["ma60"] / df["ma60"].shift(10) - 1

    df["ret5"] = df["close"] / df["close"].shift(5) - 1
    df["ret10"] = df["close"] / df["close"].shift(10) - 1
    df["ret20"] = df["close"] / df["close"].shift(20) - 1
    df["ret60"] = df["close"] / df["close"].shift(60) - 1

    df["amount_ma5"] = df["amount"].rolling(5).mean()
    df["amount_ma10"] = df["amount"].rolling(10).mean()
    df["amount_ma20"] = df["amount"].rolling(20).mean()
    df["amount_ratio_5_20"] = df["amount_ma5"] / df["amount_ma20"]

    df["turnover_ma5"] = df["turnover"].rolling(5).mean() if "turnover" in df.columns else np.nan
    df["turnover_ma20"] = df["turnover"].rolling(20).mean() if "turnover" in df.columns else np.nan

    df["high20"] = df["high"].rolling(20).max()
    df["high60"] = df["high"].rolling(60).max()
    df["is_20d_high"] = df["close"] >= df["high20"] * 0.999
    df["is_60d_high"] = df["close"] >= df["high60"] * 0.999
    df["dist_to_60d_high"] = df["close"] / df["high60"] - 1

    df["max_dd20"] = rolling_max_drawdown_np(df["close"], window=20)

    df["atr20"] = calc_atr(df, ATR_PERIOD)
    df["atr20_pct"] = df["atr20"] / df["close"]

    df["is_up_day"] = df["close"] > df["pre_close"]
    df["is_down_day"] = df["close"] < df["pre_close"]

    up_amount_20 = df["amount"].where(df["is_up_day"], 0.0).rolling(20, min_periods=20).sum()
    down_amount_20 = df["amount"].where(df["is_down_day"], 0.0).rolling(20, min_periods=20).sum()
    df["up_down_amount_ratio20"] = up_amount_20 / down_amount_20.replace(0, np.nan)

    df["is_limit_up"] = df["pct_chg"] >= 9.5
    df["is_limit_down"] = df["pct_chg"] <= -9.5
    df["limit_up_count_5"] = df["is_limit_up"].rolling(5).sum()

    intraday_range = df["high"] - df["low"]
    close_position = np.where(
        intraday_range > 0,
        (df["close"] - df["low"]) / intraday_range,
        np.nan,
    )

    df["heavy_bearish_candle"] = (
        (df["pct_chg"] <= -5)
        & (df["amount"] > df["amount_ma20"] * 1.5)
        & (close_position <= 0.35)
    )

    df["short_term_strong_raw"] = (
        (df["close"] > df["ma5"])
        & (df["ma5"] > df["ma10"])
        & (df["ma10"] > df["ma20"])
        & (df["ma5_slope_3"] > 0)
        & (df["ma10_slope_5"] > 0)
    )

    df["low5"] = df["low"].rolling(5).min()
    df["pullback_to_ma10"] = (df["low5"] / df["ma10"] - 1).abs() <= 0.035
    df["pullback_to_ma20"] = (df["low5"] / df["ma20"] - 1).abs() <= 0.045
    df["reclaim_ma10"] = (df["close"] > df["ma10"]) & (df["close"].shift(1) <= df["ma10"].shift(1))

    df["pullback_rebound_raw"] = (
        (df["ma20"] > df["ma60"])
        & (df["ma60"] > df["ma120"])
        & (df["close"] > df["ma60"])
        & (df["pullback_to_ma10"] | df["pullback_to_ma20"])
        & ((df["close"] > df["ma10"]) | df["reclaim_ma10"])
        & (df["amount_ratio_5_20"] >= 0.8)
    )

    return df


# =========================================================
# 4. 市场状态、打分、分类
# =========================================================

def get_market_state(signal_date: pd.Timestamp, bench_df: pd.DataFrame) -> Dict:
    signal_date = pd.Timestamp(signal_date).normalize()
    tmp = bench_df[bench_df["date"] <= signal_date].copy()

    if len(tmp) < 130:
        return {
            "market_ok": False,
            "market_state": "基准数据不足",
            "bench_ret20": np.nan,
            "bench_ret60": np.nan,
        }

    latest = tmp.iloc[-1]

    bench_ret20 = latest.get("ret20", np.nan)
    bench_ret60 = latest.get("ret60", np.nan)

    condition_1 = latest["close"] > latest["ma60"]
    condition_2 = latest["ma20"] > latest["ma60"]
    condition_3 = latest["ma60_slope_10"] > 0
    condition_4 = bench_ret20 > 0

    market_ok = bool(condition_1 and condition_2 and condition_3 and condition_4)

    if market_ok:
        state = "市场趋势良好"
    elif latest["close"] > latest["ma120"]:
        state = "市场中性偏强"
    else:
        state = "市场偏弱"

    return {
        "market_ok": market_ok,
        "market_state": state,
        "bench_ret20": bench_ret20,
        "bench_ret60": bench_ret60,
        "bench_close": latest["close"],
        "bench_ma20": latest["ma20"],
        "bench_ma60": latest["ma60"],
        "bench_ma120": latest["ma120"],
    }


def is_overheat(row: pd.Series) -> bool:
    if not pd.isna(row.get("ret10", np.nan)) and row["ret10"] > 0.35:
        return True
    if not pd.isna(row.get("ret20", np.nan)) and row["ret20"] > 0.60:
        return True
    if not pd.isna(row.get("limit_up_count_5", np.nan)) and row["limit_up_count_5"] >= 3:
        return True
    return False


def calc_score(row: pd.Series, bench_ret20: float, bench_ret60: float) -> int:
    score = 0

    if row["close"] > row["ma20"]:
        score += 5
    if row["ma20"] > row["ma60"]:
        score += 8
    if row["ma60"] > row["ma120"]:
        score += 7
    if row["ma20_slope_5"] > 0:
        score += 5
    if row["ma60_slope_10"] > 0:
        score += 5

    if row["close"] > row["ma5"]:
        score += 2
    if row["ma5"] > row["ma10"]:
        score += 2
    if row["ma10"] > row["ma20"]:
        score += 2
    if row["ma5_slope_3"] > 0 and row["ma10_slope_5"] > 0:
        score += 2

    if row.get("short_term_strong", False):
        score += 3

    if row.get("pullback_rebound", False):
        score += 3

    if row["ret20_rank_pct"] >= 0.70:
        score += 8
    if row["ret60_rank_pct"] >= 0.70:
        score += 8

    if row.get("industry_strong", True):
        score += 3

    if not pd.isna(bench_ret20):
        if row["ret20"] > bench_ret20 + 0.05:
            score += 5
        elif row["ret20"] > bench_ret20:
            score += 3

    if not pd.isna(bench_ret60):
        if row["ret60"] > bench_ret60 + 0.10:
            score += 4
        elif row["ret60"] > bench_ret60:
            score += 2

    ratio = row["amount_ratio_5_20"]

    if 1.2 <= ratio <= 4:
        score += 6
    elif 1.0 <= ratio < 1.2:
        score += 3

    if row["up_down_amount_ratio20"] > 1.2:
        score += 4
    elif row["up_down_amount_ratio20"] > 1.0:
        score += 2

    if ratio < 4:
        score += 2

    if row["is_20d_high"]:
        score += 4
    if row["is_60d_high"]:
        score += 6

    if row["dist_to_60d_high"] >= -0.03:
        score += 3
    elif row["dist_to_60d_high"] >= -0.05:
        score += 2

    if row["close"] > row["ma20"] and row["dist_to_60d_high"] >= -0.05:
        score += 2

    max_dd20 = row["max_dd20"]

    if not pd.isna(max_dd20):
        if max_dd20 > -0.10:
            score += 5
        elif max_dd20 > -0.15:
            score += 3
        elif max_dd20 > -0.20:
            score += 1

    atr20_pct = row.get("atr20_pct", np.nan)

    if pd.isna(atr20_pct):
        score += 1
    else:
        if atr20_pct <= 0.035:
            score += 3
        elif atr20_pct <= 0.055:
            score += 2
        elif atr20_pct <= 0.08:
            score += 1

    if not is_overheat(row):
        score += 4

    if not row["heavy_bearish_candle"]:
        score += 3

    turnover_ma5 = row["turnover_ma5"]

    if pd.isna(turnover_ma5):
        score += 2
    else:
        if turnover_ma5 < 15:
            score += 3
        elif turnover_ma5 < 25:
            score += 1

    if MARKET_FILTER_ENABLED:
        if row.get("market_ok", False):
            score += 3
        else:
            score -= 5

    return int(max(0, min(score, 100)))


def classify_status(row: pd.Series) -> str:
    if not row["basic_liquid"]:
        return "流动性不足"

    if MARKET_FILTER_ENABLED and not row.get("market_ok", False):
        return "市场环境弱观察"

    if row["close"] < row["ma60"]:
        return "跌破MA60剔除"

    if row["close"] < row["ma20"]:
        return "跌破MA20观察"

    if row["heavy_bearish_candle"]:
        return "放量长阴观察"

    if is_overheat(row):
        return "短期过热"

    if bool(row.get("trend_confirmed_by_score", False)):
        return "短中线转强池"

    if (
        row.get("short_term_strong", False)
        and row.get("basic_liquid", False)
        and row.get("trend_basic", False)
        and row.get("risk_ok", False)
        and row.get("short_term_ok", False)
        and row.get("ret20_rank_pct", 0) >= 0.70
        and row.get("score", 0) >= MIN_SCORE
    ):
        return "短线强势池"

    if row.get("pullback_rebound", False) and row["score"] >= MIN_SCORE:
        return "回踩修复观察"

    if row["score"] >= STRONG_SCORE and row["candidate"]:
        return "强趋势池"

    if row["score"] >= MIN_SCORE and row["candidate"]:
        return "趋势观察池"

    if (
        row["close"] > row["ma60"]
        and row["ma20"] > row["ma60"]
        and abs(row["close"] / row["ma20"] - 1) <= 0.05
        and row["ret60_rank_pct"] >= 0.60
    ):
        return "回踩观察"

    return "剔除"


def get_status_order() -> Dict[str, int]:
    return {
        "短中线转强池": 1,
        "短线强势池": 2,
        "强趋势池": 3,
        "趋势观察池": 4,
        "回踩修复观察": 5,
        "回踩观察": 6,
        "跌破MA20观察": 7,
        "放量长阴观察": 8,
        "短期过热": 9,
        "市场环境弱观察": 10,
        "跌破MA60剔除": 11,
        "流动性不足": 12,
        "剔除": 13,
    }


# =========================================================
# 5. 数据加载
# =========================================================

def fetch_benchmark_data() -> pd.DataFrame:
    if not os.path.exists(BENCHMARK_CACHE_FILE):
        raise FileNotFoundError(
            f"基准缓存文件不存在：{BENCHMARK_CACHE_FILE}\n"
            f"请先准备好沪深300缓存文件，并放入 data 目录。"
        )

    df = pd.read_csv(BENCHMARK_CACHE_FILE)
    df = standardize_local_df(df, code="沪深300")

    if df is None or df.empty:
        raise ValueError(f"基准缓存文件为空或字段异常：{BENCHMARK_CACHE_FILE}")

    df = add_indicators(df)

    return df


def load_stock_history_from_local(
    code: str,
    data_dir: str = "data",
    min_bars: int = 130,
) -> Optional[pd.DataFrame]:
    code = normalize_code(code)
    cache_path = os.path.join(data_dir, f"{code}.csv")

    if not os.path.exists(cache_path):
        return None

    try:
        df = pd.read_csv(cache_path)
        df = standardize_local_df(df, code=code)

        if df is None or df.empty:
            return None

        if len(df) < min_bars:
            return None

        return df

    except Exception as e:
        print(f"{code} 读取本地数据失败：{e}")
        return None


def preload_all_stock_data(codes: List[str]) -> Dict[str, pd.DataFrame]:
    data = {}

    print("正在从本地 data 目录读取股票历史数据...")
    print(f"数据目录: {CACHE_DIR}")
    print("注意：本步骤不会下载股票历史行情。")

    missing_count = 0
    insufficient_count = 0
    invalid_count = 0

    for code in tqdm(codes, desc="股票数据读取"):
        code = normalize_code(code)

        try:
            cache_path = os.path.join(CACHE_DIR, f"{code}.csv")

            if not os.path.exists(cache_path):
                missing_count += 1
                continue

            df = load_stock_history_from_local(
                code=code,
                data_dir=CACHE_DIR,
                min_bars=MIN_BARS,
            )

            if df is None or df.empty:
                insufficient_count += 1
                continue

            required_price_cols = ["open", "high", "low", "close", "amount"]
            missing_cols = [c for c in required_price_cols if c not in df.columns]

            if missing_cols:
                invalid_count += 1
                print(f"{code} 缺少字段 {missing_cols}，跳过。")
                continue

            df = add_indicators(df)
            data[code] = df

        except Exception as e:
            invalid_count += 1
            print(f"{code} 读取或处理失败：{e}")
            continue

    print(f"成功加载 {len(data)} 只股票数据。")
    print(f"本地文件不存在数量: {missing_count}")
    print(f"数据不足或为空数量: {insufficient_count}")
    print(f"字段异常或处理失败数量: {invalid_count}")

    return data


# =========================================================
# 6. 股票池构建
# =========================================================

def resolve_signal_date(
    requested_date: str,
    bench_df: pd.DataFrame,
    use_prev_trading_day: bool = False,
) -> pd.Timestamp:
    dt = pd.Timestamp(requested_date).normalize()
    trading_dates = bench_df["date"].drop_duplicates().sort_values().reset_index(drop=True)

    if (trading_dates == dt).any():
        return dt

    if not use_prev_trading_day:
        raise ValueError(
            f"{requested_date} 不是基准交易日。"
            f"如果想自动使用前一个交易日，请加参数 --use-prev-trading-day"
        )

    prev_dates = trading_dates[trading_dates < dt]

    if prev_dates.empty:
        raise ValueError(f"{requested_date} 之前没有可用交易日。")

    resolved = prev_dates.iloc[-1]
    print(f"输入日期 {requested_date} 不是交易日，已自动使用前一个交易日：{resolved.date()}")

    return resolved


def build_pool_on_date(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Optional[Dict[str, str]] = None,
    industry_map: Optional[Dict[str, str]] = None,
    classify: bool = True,
) -> pd.DataFrame:
    if code_name_map is None:
        code_name_map = {}

    if industry_map is None:
        industry_map = {}

    signal_date = pd.Timestamp(signal_date).normalize()

    market_state = get_market_state(signal_date, bench_df)
    bench_ret20 = market_state.get("bench_ret20", np.nan)
    bench_ret60 = market_state.get("bench_ret60", np.nan)

    required_indicator_cols = [
        "close",
        "ma5", "ma10", "ma20", "ma60", "ma120",
        "ma5_slope_3", "ma10_slope_5",
        "ma20_slope_5", "ma60_slope_10",
        "ret5", "ret10", "ret20", "ret60",
        "amount_ma20", "amount_ratio_5_20",
        "high20", "high60", "dist_to_60d_high",
        "max_dd20", "up_down_amount_ratio20",
        "atr20", "atr20_pct",
    ]

    rows = []

    for code, df in stock_data.items():
        if df is None or df.empty:
            continue

        df_slice = df[df["date"] <= signal_date]

        if len(df_slice) < MIN_BARS:
            continue

        latest = df_slice.iloc[-1]

        if pd.Timestamp(latest["date"]).normalize() != signal_date:
            continue

        has_nan = False

        for col in required_indicator_cols:
            if col not in latest.index or pd.isna(latest[col]):
                has_nan = True
                break

        if has_nan:
            continue

        code = normalize_code(code)
        industry = industry_map.get(code, "")

        rows.append({
            "code": code,
            "name": code_name_map.get(code, ""),
            "industry": industry,
            "signal_date": signal_date.strftime("%Y-%m-%d"),

            "close": latest["close"],
            "ma5": latest["ma5"],
            "ma10": latest["ma10"],
            "ma20": latest["ma20"],
            "ma60": latest["ma60"],
            "ma120": latest["ma120"],

            "ma5_slope_3": latest["ma5_slope_3"],
            "ma10_slope_5": latest["ma10_slope_5"],
            "ma20_slope_5": latest["ma20_slope_5"],
            "ma60_slope_10": latest["ma60_slope_10"],

            "ret5": latest["ret5"],
            "ret10": latest["ret10"],
            "ret20": latest["ret20"],
            "ret60": latest["ret60"],

            "amount_ma5": latest["amount_ma5"],
            "amount_ma10": latest["amount_ma10"],
            "amount_ma20": latest["amount_ma20"],
            "amount_ratio_5_20": latest["amount_ratio_5_20"],

            "turnover_ma5": latest.get("turnover_ma5", np.nan),
            "turnover_ma20": latest.get("turnover_ma20", np.nan),

            "high20": latest["high20"],
            "high60": latest["high60"],
            "is_20d_high": bool(latest["is_20d_high"]),
            "is_60d_high": bool(latest["is_60d_high"]),
            "dist_to_60d_high": latest["dist_to_60d_high"],

            "max_dd20": latest["max_dd20"],
            "atr20": latest["atr20"],
            "atr20_pct": latest["atr20_pct"],

            "up_down_amount_ratio20": latest["up_down_amount_ratio20"],

            "limit_up_count_5": latest.get("limit_up_count_5", np.nan),
            "heavy_bearish_candle": bool(latest.get("heavy_bearish_candle", False)),

            "short_term_strong": bool(latest.get("short_term_strong_raw", False)),
            "pullback_rebound": bool(latest.get("pullback_rebound_raw", False)),

            "bench_ret20": bench_ret20,
            "bench_ret60": bench_ret60,

            "market_ok": market_state.get("market_ok", False),
            "market_state": market_state.get("market_state", ""),
        })

    if not rows:
        return pd.DataFrame()

    df_pool = pd.DataFrame(rows)

    df_pool["ret20_rank_pct"] = df_pool["ret20"].rank(pct=True)
    df_pool["ret60_rank_pct"] = df_pool["ret60"].rank(pct=True)

    if INDUSTRY_FILTER_ENABLED and "industry" in df_pool.columns and df_pool["industry"].astype(str).str.len().gt(0).any():
        industry_strength = (
            df_pool[df_pool["industry"].astype(str).str.len() > 0]
            .groupby("industry")
            .agg(
                industry_ret20=("ret20", "mean"),
                industry_ret60=("ret60", "mean"),
                industry_count=("code", "count"),
            )
            .reset_index()
        )

        industry_strength["industry_ret20_rank_pct"] = industry_strength["industry_ret20"].rank(pct=True)
        industry_strength["industry_ret60_rank_pct"] = industry_strength["industry_ret60"].rank(pct=True)

        df_pool = df_pool.merge(industry_strength, on="industry", how="left")

        df_pool["industry_strong"] = (
            (df_pool["industry_ret20_rank_pct"] >= 0.60)
            & (df_pool["industry_ret60_rank_pct"] >= 0.60)
        )

        df_pool["industry_strong"] = df_pool["industry_strong"].fillna(True)
    else:
        df_pool["industry_ret20"] = np.nan
        df_pool["industry_ret60"] = np.nan
        df_pool["industry_ret20_rank_pct"] = np.nan
        df_pool["industry_ret60_rank_pct"] = np.nan
        df_pool["industry_count"] = np.nan
        df_pool["industry_strong"] = True

    df_pool["basic_liquid"] = df_pool["amount_ma20"] >= MIN_AVG_AMOUNT_20

    df_pool["trend_basic"] = (
        (df_pool["close"] > df_pool["ma20"])
        & (df_pool["ma20"] > df_pool["ma60"])
        & (df_pool["ma60"] > df_pool["ma120"])
        & (df_pool["ma20_slope_5"] > 0)
        & (df_pool["ma60_slope_10"] > 0)
    )

    df_pool["relative_strength"] = (
        (df_pool["ret20_rank_pct"] >= 0.70)
        & (df_pool["ret60_rank_pct"] >= 0.70)
    )

    df_pool["volume_ok"] = (
        (df_pool["amount_ratio_5_20"] >= 1.2)
        & (df_pool["amount_ratio_5_20"] <= 4)
    )

    df_pool["near_breakout"] = (
        (df_pool["dist_to_60d_high"] >= -0.05)
        | df_pool["is_20d_high"]
        | df_pool["is_60d_high"]
    )

    df_pool["risk_ok"] = (
        (df_pool["close"] > df_pool["ma20"])
        & (df_pool["max_dd20"] > -0.20)
        & (~df_pool["heavy_bearish_candle"])
        & (df_pool["atr20_pct"] <= 0.10)
    )

    df_pool["short_term_ok"] = (
        df_pool["short_term_strong"]
        | df_pool["pullback_rebound"]
        | (
            (df_pool["close"] > df_pool["ma10"])
            & (df_pool["ma10"] > df_pool["ma20"])
        )
    )

    df_pool["candidate"] = (
        df_pool["basic_liquid"]
        & df_pool["trend_basic"]
        & df_pool["relative_strength"]
        & df_pool["volume_ok"]
        & df_pool["near_breakout"]
        & df_pool["risk_ok"]
        & df_pool["industry_strong"]
    )

    if MARKET_FILTER_ENABLED:
        df_pool["candidate"] = df_pool["candidate"] & df_pool["market_ok"]

    df_pool["score"] = df_pool.apply(
        lambda row: calc_score(row, bench_ret20, bench_ret60),
        axis=1,
    )

    df_pool["trend_confirmed_by_score"] = False

    df_pool["status"] = df_pool.apply(classify_status, axis=1)

    df_pool = sort_pool_df(df_pool)

    return df_pool


def sort_pool_df(df_pool: pd.DataFrame) -> pd.DataFrame:
    if df_pool is None or df_pool.empty:
        return pd.DataFrame()

    status_order = get_status_order()
    df_pool = df_pool.copy()
    df_pool["_order"] = df_pool["status"].map(status_order).fillna(99)
    df_pool = df_pool.sort_values(
        ["_order", "score"],
        ascending=[True, False],
    ).drop(columns=["_order"]).reset_index(drop=True)

    return df_pool


# =========================================================
# 7. 连续多日评分趋势
# =========================================================

def classify_score_trend(g: pd.DataFrame) -> pd.Series:
    g = g.sort_values("signal_date").copy()

    scores = g["score"].dropna().astype(float).tolist()
    statuses = g["status"].astype(str).tolist()

    if len(scores) < 4:
        return pd.Series({
            "score_trend_status": "样本不足",
            "score_now": np.nan,
            "score_3d_ago": np.nan,
            "score_5d_ago": np.nan,
            "score_change_3d": np.nan,
            "score_change_5d": np.nan,
            "score_ma3": np.nan,
            "score_ma5": np.nan,
            "score_min5": np.nan,
            "score_max5": np.nan,
            "status_5d_ago": "",
            "status_changed": False,
        })

    last_scores = scores[-SCORE_TREND_WINDOW:] if len(scores) >= SCORE_TREND_WINDOW else scores

    score_now = last_scores[-1]
    score_first = last_scores[0]

    score_3d_ago = scores[-4] if len(scores) >= 4 else np.nan
    score_5d_ago = scores[-6] if len(scores) >= 6 else score_first

    score_change_3d = score_now - score_3d_ago if not pd.isna(score_3d_ago) else np.nan
    score_change_5d = score_now - score_first

    score_ma3 = np.mean(last_scores[-3:]) if len(last_scores) >= 3 else np.mean(last_scores)
    score_ma5 = np.mean(last_scores)
    score_min5 = np.min(last_scores)
    score_max5 = np.max(last_scores)

    status_5d_ago = statuses[-6] if len(statuses) >= 6 else statuses[0]
    status_now = statuses[-1]
    status_changed = status_now != status_5d_ago

    if score_now >= 85 and score_min5 >= 75 and score_ma3 >= 85:
        status = "强趋势延续"
    elif score_now >= 70 and score_change_5d >= 20 and score_ma3 > score_ma5:
        status = "趋势转强"
    elif score_now >= 60 and score_change_5d >= 15:
        status = "趋势修复"
    elif score_first >= 85 and score_now <= 70 and score_ma3 < score_ma5:
        status = "高位转弱"
    elif score_now < 50 and score_ma5 < 50:
        status = "弱势延续"
    else:
        status = "震荡无趋势"

    return pd.Series({
        "score_trend_status": status,
        "score_now": score_now,
        "score_3d_ago": score_3d_ago,
        "score_5d_ago": score_5d_ago,
        "score_change_3d": score_change_3d,
        "score_change_5d": score_change_5d,
        "score_ma3": score_ma3,
        "score_ma5": score_ma5,
        "score_min5": score_min5,
        "score_max5": score_max5,
        "status_5d_ago": status_5d_ago,
        "status_changed": status_changed,
    })


def build_score_history(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Dict[str, str],
    industry_map: Dict[str, str],
    lookback_days: int = 10,
) -> pd.DataFrame:
    signal_date = pd.Timestamp(signal_date).normalize()

    trading_dates = (
        bench_df[bench_df["date"] <= signal_date]["date"]
        .drop_duplicates()
        .sort_values()
    )

    dates = trading_dates.tail(lookback_days).tolist()

    frames = []

    print(f"\n正在构建最近 {lookback_days} 个交易日评分历史...")

    for d in tqdm(dates, desc="评分历史"):
        pool = build_pool_on_date(
            signal_date=d,
            stock_data=stock_data,
            bench_df=bench_df,
            code_name_map=code_name_map,
            industry_map=industry_map,
            classify=True,
        )

        if pool is None or pool.empty:
            continue

        keep_cols = [
            "code", "name", "industry", "signal_date", "score", "status",
            "ret20_rank_pct", "ret60_rank_pct",
            "close", "ma5", "ma10", "ma20", "ma60",
            "ma20_slope_5",
            "short_term_strong", "pullback_rebound",
            "trend_basic", "risk_ok", "market_ok",
        ]
        keep_cols = [c for c in keep_cols if c in pool.columns]
        tmp = pool[keep_cols].copy()
        frames.append(tmp)

    if not frames:
        return pd.DataFrame()

    out = pd.concat(frames, ignore_index=True)
    out["signal_date"] = pd.to_datetime(out["signal_date"], errors="coerce").dt.strftime("%Y-%m-%d")
    return out


def add_score_trend_to_pool(
    pool_df: pd.DataFrame,
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Dict[str, str],
    industry_map: Dict[str, str],
    lookback_days: int = 10,
) -> pd.DataFrame:
    if pool_df is None or pool_df.empty:
        return pd.DataFrame()

    score_history_df = build_score_history(
        signal_date=signal_date,
        stock_data=stock_data,
        bench_df=bench_df,
        code_name_map=code_name_map,
        industry_map=industry_map,
        lookback_days=lookback_days,
    )

    if score_history_df is None or score_history_df.empty:
        pool_df = pool_df.copy()
        pool_df["score_trend_status"] = "样本不足"
        pool_df["score_change_5d"] = np.nan
        pool_df["score_ma3"] = np.nan
        pool_df["score_ma5"] = np.nan
        pool_df["trend_confirmed_by_score"] = False
        pool_df["status_changed"] = False
        return pool_df

    score_trend_df = (
        score_history_df
        .groupby("code")
        .apply(classify_score_trend)
        .reset_index()
    )

    out = pool_df.merge(score_trend_df, on="code", how="left")

    out["trend_confirmed_by_score"] = (
        (out["score"] >= 75)
        & (out["score_change_5d"] >= 15)
        & (out["close"] > out["ma20"])
        & (out["ma20"] > out["ma60"])
        & (out["ma20_slope_5"] > 0)
        & (out["ret20_rank_pct"] >= 0.70)
        & ((out["short_term_strong"]) | (out["pullback_rebound"]))
        & (out["risk_ok"])
    )

    out["status"] = out.apply(classify_status, axis=1)
    out = sort_pool_df(out)

    return out


# =========================================================
# 8. 统计和 Excel
# =========================================================

def build_classification_summary(pool_df: pd.DataFrame) -> pd.DataFrame:
    if pool_df is None or pool_df.empty:
        return pd.DataFrame()

    total = len(pool_df)

    agg_dict = dict(
        数量=("code", "count"),
        平均评分=("score", "mean"),
        最高评分=("score", "max"),
        平均20日涨幅=("ret20", "mean"),
        平均60日涨幅=("ret60", "mean"),
        平均20日成交额=("amount_ma20", "mean"),
        平均量能比=("amount_ratio_5_20", "mean"),
        平均20日最大回撤=("max_dd20", "mean"),
        平均ATR比例=("atr20_pct", "mean"),
    )

    if "score_change_5d" in pool_df.columns:
        agg_dict["平均5日评分变化"] = ("score_change_5d", "mean")

    summary = (
        pool_df
        .groupby("status")
        .agg(**agg_dict)
        .reset_index()
        .rename(columns={"status": "分类"})
    )

    summary["占比"] = summary["数量"] / total

    status_order = get_status_order()
    summary["排序"] = summary["分类"].map(status_order).fillna(99)
    summary = summary.sort_values(["排序", "数量"], ascending=[True, False])
    summary = summary.drop(columns=["排序"]).reset_index(drop=True)

    return summary


def make_classification_readable(pool_df: pd.DataFrame) -> pd.DataFrame:
    if pool_df is None or pool_df.empty:
        return pd.DataFrame()

    df = pool_df.copy()

    keep_cols = [
        "code", "name", "industry", "signal_date", "status", "score",
        "score_trend_status", "score_now",
        "score_3d_ago", "score_5d_ago",
        "score_change_3d", "score_change_5d",
        "score_ma3", "score_ma5", "score_min5", "score_max5",
        "status_5d_ago", "status_changed", "trend_confirmed_by_score",
        "market_state", "market_ok",
        "close",
        "ma5", "ma10", "ma20", "ma60", "ma120",
        "ma5_slope_3", "ma10_slope_5", "ma20_slope_5", "ma60_slope_10",
        "ret5", "ret10", "ret20", "ret60",
        "ret20_rank_pct", "ret60_rank_pct",
        "industry_ret20", "industry_ret60",
        "industry_ret20_rank_pct", "industry_ret60_rank_pct",
        "industry_strong",
        "amount_ma5", "amount_ma10", "amount_ma20", "amount_ratio_5_20",
        "turnover_ma5", "turnover_ma20",
        "is_20d_high", "is_60d_high", "dist_to_60d_high",
        "max_dd20", "atr20", "atr20_pct",
        "up_down_amount_ratio20",
        "limit_up_count_5", "heavy_bearish_candle",
        "short_term_strong", "pullback_rebound",
        "bench_ret20", "bench_ret60",
        "basic_liquid", "trend_basic", "relative_strength",
        "volume_ok", "near_breakout", "risk_ok", "short_term_ok",
        "candidate",
    ]

    keep_cols = [c for c in keep_cols if c in df.columns]
    out = df[keep_cols].copy()

    rename_map = {
        "code": "代码",
        "name": "名称",
        "industry": "行业",
        "signal_date": "信号日",
        "status": "分类",
        "score": "评分",

        "score_trend_status": "评分趋势状态",
        "score_now": "当前评分",
        "score_3d_ago": "3日前评分",
        "score_5d_ago": "5日前评分",
        "score_change_3d": "近3日评分变化",
        "score_change_5d": "近5日评分变化",
        "score_ma3": "近3日平均评分",
        "score_ma5": "近5日平均评分",
        "score_min5": "近5日最低评分",
        "score_max5": "近5日最高评分",
        "status_5d_ago": "5日前分类",
        "status_changed": "分类是否变化",
        "trend_confirmed_by_score": "评分确认转强",

        "market_state": "市场状态",
        "market_ok": "大盘环境通过",
        "close": "收盘价",
        "ma5": "MA5",
        "ma10": "MA10",
        "ma20": "MA20",
        "ma60": "MA60",
        "ma120": "MA120",
        "ma5_slope_3": "MA5近3日斜率",
        "ma10_slope_5": "MA10近5日斜率",
        "ma20_slope_5": "MA20近5日斜率",
        "ma60_slope_10": "MA60近10日斜率",
        "ret5": "近5日涨幅",
        "ret10": "近10日涨幅",
        "ret20": "近20日涨幅",
        "ret60": "近60日涨幅",
        "ret20_rank_pct": "20日强度排名百分位",
        "ret60_rank_pct": "60日强度排名百分位",
        "industry_ret20": "行业20日涨幅",
        "industry_ret60": "行业60日涨幅",
        "industry_ret20_rank_pct": "行业20日强度百分位",
        "industry_ret60_rank_pct": "行业60日强度百分位",
        "industry_strong": "行业强度达标",
        "amount_ma5": "5日均成交额",
        "amount_ma10": "10日均成交额",
        "amount_ma20": "20日均成交额",
        "amount_ratio_5_20": "量能比5/20",
        "turnover_ma5": "5日均换手率",
        "turnover_ma20": "20日均换手率",
        "is_20d_high": "是否接近20日新高",
        "is_60d_high": "是否近60日新高",
        "dist_to_60d_high": "距60日高点",
        "max_dd20": "近20日最大回撤",
        "atr20": "ATR20",
        "atr20_pct": "ATR20比例",
        "up_down_amount_ratio20": "近20日上涨/下跌成交额比",
        "limit_up_count_5": "近5日涨停次数",
        "heavy_bearish_candle": "是否放量长阴",
        "short_term_strong": "短线强势",
        "pullback_rebound": "回踩修复",
        "bench_ret20": "沪深300近20日涨幅",
        "bench_ret60": "沪深300近60日涨幅",
        "basic_liquid": "流动性达标",
        "trend_basic": "中期趋势达标",
        "relative_strength": "相对强度达标",
        "volume_ok": "量能达标",
        "near_breakout": "接近突破",
        "risk_ok": "风险过滤通过",
        "short_term_ok": "短线结构通过",
        "candidate": "是否候选",
    }

    out = out.rename(columns=rename_map)

    status_order = get_status_order()

    if "分类" in out.columns:
        out["_排序"] = out["分类"].map(status_order).fillna(99)

        if "评分" in out.columns:
            out = out.sort_values(["_排序", "评分"], ascending=[True, False])
        else:
            out = out.sort_values(["_排序"], ascending=True)

        out = out.drop(columns=["_排序"]).reset_index(drop=True)

    return out


def style_worksheet(
    ws,
    freeze: str = "A2",
    percent_keywords: Optional[List[str]] = None,
    money_keywords: Optional[List[str]] = None,
    integer_keywords: Optional[List[str]] = None,
    apply_red_green: bool = True,
):
    if percent_keywords is None:
        percent_keywords = [
            "收益", "超额", "涨幅", "回撤", "胜率", "比例", "D",
            "距60日高点", "ATR", "斜率", "占比", "百分位",
        ]

    if money_keywords is None:
        money_keywords = ["成交额", "买入价", "卖出价", "收盘价", "MA", "ATR"]

    if integer_keywords is None:
        integer_keywords = ["天数", "样本数", "评分", "跑赢天数", "有效天数", "持有天数", "数量", "次数"]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    max_row = ws.max_row
    max_col = ws.max_column

    positive_fill = PatternFill("solid", fgColor="FCE4D6")
    negative_fill = PatternFill("solid", fgColor="E2F0D9")

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        header = ws.cell(row=1, column=col_idx).value
        header_str = "" if header is None else str(header)

        max_len = len(header_str)

        for row_idx in range(2, min(max_row, 200) + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None:
                max_len = max(max_len, len(str(v)))

        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 32)

        is_percent_col = any(k in header_str for k in percent_keywords)
        is_money_col = any(k in header_str for k in money_keywords)
        is_integer_col = any(k in header_str for k in integer_keywords)

        for row_idx in range(2, max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = Alignment(vertical="center")

            if is_percent_col and isinstance(cell.value, (int, float)) and not pd.isna(cell.value):
                cell.number_format = "0.00%"
            elif is_money_col and isinstance(cell.value, (int, float)) and not pd.isna(cell.value):
                cell.number_format = "#,##0.00"
            elif is_integer_col and isinstance(cell.value, (int, float)) and not pd.isna(cell.value):
                cell.number_format = "0"

        if apply_red_green and is_percent_col and max_row >= 2:
            data_range = f"{col_letter}2:{col_letter}{max_row}"

            ws.conditional_formatting.add(
                data_range,
                CellIsRule(
                    operator="greaterThan",
                    formula=["0"],
                    fill=positive_fill,
                ),
            )

            ws.conditional_formatting.add(
                data_range,
                CellIsRule(
                    operator="lessThan",
                    formula=["0"],
                    fill=negative_fill,
                ),
            )


def save_classification_excel(
    output_xlsx: str,
    signal_date: pd.Timestamp,
    pool_df: pd.DataFrame,
) -> None:
    readable_df = make_classification_readable(pool_df)
    summary_df = build_classification_summary(pool_df)

    dashboard_rows = []
    total_count = len(pool_df) if pool_df is not None else 0

    dashboard_rows.append(["信号日期", signal_date.strftime("%Y-%m-%d"), "分类所使用的交易日"])
    dashboard_rows.append(["全市场有效样本数", total_count, "满足数据长度和指标完整性的股票数量"])
    dashboard_rows.append(["最低评分参数", MIN_SCORE, "趋势观察池最低评分"])
    dashboard_rows.append(["强趋势评分参数", STRONG_SCORE, "强趋势池最低评分"])
    dashboard_rows.append(["评分趋势回看交易日", SCORE_TREND_LOOKBACK, "用于 score_trend_status"])
    dashboard_rows.append(["评分趋势窗口", SCORE_TREND_WINDOW, "默认最近5日计算 score_change_5d"])
    dashboard_rows.append(["20日均成交额门槛", MIN_AVG_AMOUNT_20, "流动性过滤"])
    dashboard_rows.append(["最少K线数量", MIN_BARS, "参与计算要求"])
    dashboard_rows.append(["大盘环境过滤", "开启" if MARKET_FILTER_ENABLED else "关闭", "沪深300趋势过滤"])
    dashboard_rows.append(["行业强度过滤", "开启" if INDUSTRY_FILTER_ENABLED else "关闭", "依赖 data/_stock_industry.csv"])

    if summary_df is not None and not summary_df.empty:
        for _, r in summary_df.iterrows():
            dashboard_rows.append([
                f"{r['分类']}数量",
                r["数量"],
                f"占比 {r['占比']:.2%}",
            ])

    dashboard_df = pd.DataFrame(dashboard_rows, columns=["指标", "数值", "说明"])

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        dashboard_df.to_excel(writer, sheet_name="看板", index=False)
        summary_df.to_excel(writer, sheet_name="分类统计", index=False)
        readable_df.to_excel(writer, sheet_name="全市场分类", index=False)

        important_statuses = list(get_status_order().keys())

        if not readable_df.empty and "分类" in readable_df.columns:
            for status in important_statuses:
                sub = readable_df[readable_df["分类"] == status].copy()

                if sub.empty:
                    continue

                sheet_name = status[:31]
                sub.to_excel(writer, sheet_name=sheet_name, index=False)

        wb = writer.book

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            if sheet_name == "看板":
                style_worksheet(ws, freeze="A2", apply_red_green=True)
                ws.column_dimensions["A"].width = 28
                ws.column_dimensions["B"].width = 20
                ws.column_dimensions["C"].width = 60
            else:
                style_worksheet(ws, freeze="A2", apply_red_green=True)


# =========================================================
# 9. 固定持有回测
# =========================================================

def get_forward_trading_dates(
    signal_date: pd.Timestamp,
    bench_df: pd.DataFrame,
    forward_days: int,
) -> List[pd.Timestamp]:
    signal_date = pd.Timestamp(signal_date).normalize()

    future_dates = bench_df.loc[
        bench_df["date"] > signal_date,
        "date"
    ].head(forward_days).tolist()

    return [pd.Timestamp(d).normalize() for d in future_dates]


def calc_forward_return_path_by_calendar(
    df: pd.DataFrame,
    forward_dates: List[pd.Timestamp],
    buy_price_field: str = BUY_PRICE_FIELD,
) -> Dict:
    if df is None or df.empty or not forward_dates:
        return {}

    df_map = df.set_index("date", drop=False)
    buy_date = forward_dates[0]

    if buy_date not in df_map.index:
        return {}

    buy_row = df_map.loc[buy_date]

    if isinstance(buy_row, pd.DataFrame):
        buy_row = buy_row.iloc[-1]

    if is_open_limit_up(buy_row):
        return {}

    raw_buy_price = get_row_price(buy_row, buy_price_field, fallback="close")
    buy_price = net_buy_price(raw_buy_price)

    if pd.isna(buy_price) or buy_price <= 0:
        return {}

    result = {
        "buy_date": buy_date.strftime("%Y-%m-%d"),
        "raw_buy_price": raw_buy_price,
        "buy_price": buy_price,
    }

    returns = []

    for i, d in enumerate(forward_dates, start=1):
        if d not in df_map.index:
            ret = np.nan
        else:
            row = df_map.loc[d]

            if isinstance(row, pd.DataFrame):
                row = row.iloc[-1]

            close_price = get_row_price(row, "close", fallback="close")
            ret = calc_net_return(close_price, raw_buy_price)

        result[f"D{i}"] = ret
        returns.append(ret)

    ret_series = pd.Series(returns).dropna()
    n = len(forward_dates)

    if ret_series.empty:
        result[f"max_ret_{n}"] = np.nan
        result[f"min_ret_{n}"] = np.nan
        result[f"max_dd_{n}"] = np.nan
        return result

    result[f"max_ret_{n}"] = ret_series.max()
    result[f"min_ret_{n}"] = ret_series.min()
    result[f"max_dd_{n}"] = max_drawdown(1 + ret_series)

    return result


def build_forward_summary(detail_df: pd.DataFrame, forward_days: int, group_col: str = "status") -> pd.DataFrame:
    if detail_df is None or detail_df.empty:
        return pd.DataFrame()

    bench_rows = detail_df[detail_df["row_type"] == "benchmark"]
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    if bench_rows.empty or stock_rows.empty:
        return pd.DataFrame()

    bench_row = bench_rows.iloc[0]
    rows = []

    for group_value, group in stock_rows.groupby(group_col):
        for h in range(1, forward_days + 1):
            col = f"D{h}"

            if col not in group.columns:
                continue

            bench_ret = bench_row[col]
            valid = group[group[col].notna()].copy()

            if valid.empty:
                continue

            ret = valid[col]
            excess = ret - bench_ret

            rows.append({
                group_col: group_value,
                "horizon": h,
                "count": len(valid),
                "avg_ret": ret.mean(),
                "median_ret": ret.median(),
                "win_rate": (ret > 0).mean(),
                "bench_ret": bench_ret,
                "avg_excess": excess.mean(),
                "outperform_rate": (excess > 0).mean(),
            })

    return pd.DataFrame(rows)


def build_overall_forward_summary(detail_df: pd.DataFrame, forward_days: int) -> pd.DataFrame:
    if detail_df is None or detail_df.empty:
        return pd.DataFrame()

    bench_rows = detail_df[detail_df["row_type"] == "benchmark"]
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    if bench_rows.empty or stock_rows.empty:
        return pd.DataFrame()

    bench_row = bench_rows.iloc[0]
    rows = []

    for h in range(1, forward_days + 1):
        col = f"D{h}"

        if col not in stock_rows.columns:
            continue

        bench_ret = bench_row[col]
        valid = stock_rows[stock_rows[col].notna()].copy()

        if valid.empty:
            continue

        ret = valid[col]
        excess = ret - bench_ret

        rows.append({
            "horizon": h,
            "count": len(valid),
            "avg_ret": ret.mean(),
            "median_ret": ret.median(),
            "win_rate": (ret > 0).mean(),
            "bench_ret": bench_ret,
            "avg_excess": excess.mean(),
            "outperform_rate": (excess > 0).mean(),
        })

    return pd.DataFrame(rows)


def rename_forward_summary_columns(df: pd.DataFrame, group_col: Optional[str] = None) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()

    out = df.copy()

    rename_map = {
        "horizon": "持有天数",
        "count": "样本数",
        "avg_ret": "平均收益",
        "median_ret": "中位数收益",
        "win_rate": "胜率",
        "bench_ret": "沪深300收益",
        "avg_excess": "平均超额",
        "outperform_rate": "跑赢沪深300比例",
        "status": "分组",
    }

    if group_col and group_col in out.columns:
        rename_map[group_col] = "分组"

    return out.rename(columns=rename_map)


def save_forward_excel(
    output_xlsx: str,
    signal_date: pd.Timestamp,
    detail_df: pd.DataFrame,
    group_summary_df: pd.DataFrame,
    overall_summary_df: pd.DataFrame,
    selected_df: pd.DataFrame,
    pool_df: pd.DataFrame,
    forward_days: int,
) -> None:
    dashboard_rows = []
    stock_rows = detail_df[detail_df["row_type"] == "stock"].copy()

    dashboard_rows.append(["信号日期", signal_date.strftime("%Y-%m-%d"), "策略信号生成日期"])
    dashboard_rows.append(["买入价格口径", BUY_PRICE_FIELD, "open=买入日开盘价，close=买入日收盘价"])
    dashboard_rows.append(["未来观察天数", forward_days, f"D1~D{forward_days}"])
    dashboard_rows.append(["入选股票数", len(stock_rows), "进入目标池且可买入的股票数量"])
    dashboard_rows.append(["评分趋势回看交易日", SCORE_TREND_LOOKBACK, "用于 score_trend_status"])

    if overall_summary_df is not None and not overall_summary_df.empty:
        for _, r in overall_summary_df[overall_summary_df["horizon"].isin([1, 3, 5, 10, 20, 30])].iterrows():
            h = int(r["horizon"])
            dashboard_rows.append([f"D{h}平均收益", r["avg_ret"], "固定持有口径，已含成本滑点"])
            dashboard_rows.append([f"D{h}胜率", r["win_rate"], "固定持有口径"])
            dashboard_rows.append([f"D{h}平均超额收益", r["avg_excess"], "固定持有平均收益 - 沪深300"])

    dashboard_df = pd.DataFrame(dashboard_rows, columns=["指标", "数值", "说明"])

    readable_pool_df = make_classification_readable(pool_df)
    selected_readable_df = make_classification_readable(selected_df)
    overall_readable = rename_forward_summary_columns(overall_summary_df)
    group_readable = rename_forward_summary_columns(group_summary_df, group_col="status")

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        dashboard_df.to_excel(writer, sheet_name="看板", index=False)
        selected_readable_df.to_excel(writer, sheet_name="入选股票", index=False)
        overall_readable.to_excel(writer, sheet_name="整体统计", index=False)
        group_readable.to_excel(writer, sheet_name="分组统计", index=False)
        detail_df.to_excel(writer, sheet_name="收益明细", index=False)
        readable_pool_df.to_excel(writer, sheet_name="全市场分类", index=False)

        wb = writer.book

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            if sheet_name == "看板":
                style_worksheet(ws, freeze="A2", apply_red_green=True)
                ws.column_dimensions["A"].width = 28
                ws.column_dimensions["B"].width = 20
                ws.column_dimensions["C"].width = 60
            else:
                style_worksheet(ws, freeze="A2", apply_red_green=True)


# =========================================================
# 10. 主分析函数
# =========================================================

def build_final_pool_with_score_trend(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Dict[str, str],
    industry_map: Dict[str, str],
    score_trend_lookback: int,
) -> pd.DataFrame:
    pool_df = build_pool_on_date(
        signal_date=signal_date,
        stock_data=stock_data,
        bench_df=bench_df,
        code_name_map=code_name_map,
        industry_map=industry_map,
        classify=True,
    )

    if pool_df is None or pool_df.empty:
        return pd.DataFrame()

    pool_df = add_score_trend_to_pool(
        pool_df=pool_df,
        signal_date=signal_date,
        stock_data=stock_data,
        bench_df=bench_df,
        code_name_map=code_name_map,
        industry_map=industry_map,
        lookback_days=score_trend_lookback,
    )

    return pool_df


def analyze_pool_classification(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Dict[str, str],
    industry_map: Dict[str, str],
    score_trend_lookback: int,
) -> None:
    signal_date = pd.Timestamp(signal_date).normalize()

    print("\n========== 指定日期股票分类 ==========")
    print(f"信号日期: {signal_date.date()}")
    print("功能: 仅分类，不计算未来收益")
    print("====================================")

    pool_df = build_final_pool_with_score_trend(
        signal_date=signal_date,
        stock_data=stock_data,
        bench_df=bench_df,
        code_name_map=code_name_map,
        industry_map=industry_map,
        score_trend_lookback=score_trend_lookback,
    )

    if pool_df is None or pool_df.empty:
        print("指定日期未能构建股票分类结果。")
        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    output_xlsx = os.path.join(
        OUTPUT_DIR,
        f"{CLASSIFICATION_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.xlsx",
    )

    output_csv = os.path.join(
        OUTPUT_DIR,
        f"{CLASSIFICATION_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.csv",
    )

    try:
        save_classification_excel(
            output_xlsx=output_xlsx,
            signal_date=signal_date,
            pool_df=pool_df,
        )
        print(f"\n分类结果 Excel 已保存到：{output_xlsx}")
    except Exception as e:
        print(f"保存分类 Excel 失败：{e}")

    try:
        pool_df.to_csv(output_csv, index=False, encoding="utf-8-sig")
        print(f"分类结果 CSV 已保存到：{output_csv}")
    except Exception as e:
        print(f"保存分类 CSV 失败：{e}")

    summary_df = build_classification_summary(pool_df)

    if not summary_df.empty:
        print("\n========== 分类统计 ==========")

        formatters = {
            "占比": "{:.2%}".format,
            "平均评分": "{:.2f}".format,
            "最高评分": "{:.2f}".format,
            "平均20日涨幅": "{:.2%}".format,
            "平均60日涨幅": "{:.2%}".format,
            "平均20日成交额": "{:,.0f}".format,
            "平均量能比": "{:.2f}".format,
            "平均20日最大回撤": "{:.2%}".format,
            "平均ATR比例": "{:.2%}".format,
        }

        if "平均5日评分变化" in summary_df.columns:
            formatters["平均5日评分变化"] = "{:.2f}".format

        print(summary_df.to_string(index=False, formatters=formatters))

    selected_df = pool_df[pool_df["status"].isin([
        "短中线转强池",
        "短线强势池",
        "强趋势池",
        "趋势观察池",
        "回踩修复观察",
    ])].copy()

    selected_df = selected_df.sort_values("score", ascending=False).reset_index(drop=True)

    if not selected_df.empty:
        preview_cols = [
            "code", "name", "industry", "status", "score",
            "score_trend_status", "score_change_5d",
            "ret20", "ret60", "amount_ma20",
            "amount_ratio_5_20", "max_dd20", "atr20_pct",
            "short_term_strong", "pullback_rebound",
            "trend_confirmed_by_score",
            "dist_to_60d_high",
        ]

        preview_cols = [c for c in preview_cols if c in selected_df.columns]

        print("\n========== 核心股票池预览 ==========")
        print(selected_df[preview_cols].head(30).to_string(index=False, formatters={
            "score_change_5d": "{:.2f}".format,
            "ret20": "{:.2%}".format,
            "ret60": "{:.2%}".format,
            "amount_ma20": "{:,.0f}".format,
            "amount_ratio_5_20": "{:.2f}".format,
            "max_dd20": "{:.2%}".format,
            "atr20_pct": "{:.2%}".format,
            "dist_to_60d_high": "{:.2%}".format,
        }))
    else:
        print("\n指定日期没有股票进入核心股票池。")


def analyze_pool_forward_returns(
    signal_date: pd.Timestamp,
    stock_data: Dict[str, pd.DataFrame],
    bench_df: pd.DataFrame,
    code_name_map: Dict[str, str],
    industry_map: Dict[str, str],
    forward_days: int,
    statuses: List[str],
    score_trend_lookback: int,
) -> None:
    signal_date = pd.Timestamp(signal_date).normalize()

    print("\n========== 指定日期股票池未来走势分析 ==========")
    print(f"信号日期: {signal_date.date()}")
    print(f"分析池: {statuses}")
    print(f"未来交易日数: {forward_days}")
    print("==============================================")

    pool_df = build_final_pool_with_score_trend(
        signal_date=signal_date,
        stock_data=stock_data,
        bench_df=bench_df,
        code_name_map=code_name_map,
        industry_map=industry_map,
        score_trend_lookback=score_trend_lookback,
    )

    if pool_df.empty:
        print("指定日期未能构建股票池。")
        return

    selected_df = pool_df[pool_df["status"].isin(statuses)].copy()
    selected_df = selected_df.sort_values("score", ascending=False).reset_index(drop=True)

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    if selected_df.empty:
        print(f"{signal_date.date()} 没有股票进入 {statuses}。")

        output_all_pool = os.path.join(
            OUTPUT_DIR,
            f"{FORWARD_ANALYSIS_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}_all_pool.csv",
        )

        pool_df.to_csv(output_all_pool, index=False, encoding="utf-8-sig")
        print(f"已保存当日全市场分类到：{output_all_pool}")

        return

    print(f"入选股票数量: {len(selected_df)}")

    preview_cols = [
        "code", "name", "industry", "status", "score",
        "score_trend_status", "score_change_5d",
        "ret20", "ret60",
        "amount_ma20", "amount_ratio_5_20", "atr20_pct",
        "short_term_strong", "pullback_rebound",
        "trend_confirmed_by_score",
    ]

    preview_cols = [c for c in preview_cols if c in selected_df.columns]

    print("\n入选股票预览：")
    print(selected_df[preview_cols].head(20).to_string(index=False))

    forward_dates = get_forward_trading_dates(
        signal_date=signal_date,
        bench_df=bench_df,
        forward_days=forward_days,
    )

    if len(forward_dates) < forward_days:
        print(f"基准未来交易日不足：需要 {forward_days} 天，实际只有 {len(forward_dates)} 天。")
        print("请降低 --days，或者选择更早的信号日期。")
        return

    bench_path = calc_forward_return_path_by_calendar(
        df=bench_df,
        forward_dates=forward_dates,
        buy_price_field=BUY_PRICE_FIELD,
    )

    if not bench_path:
        print("基准未来走势计算失败。")
        return

    detail_rows = []

    bench_row = {
        "row_type": "benchmark",
        "code": "沪深300",
        "name": "沪深300",
        "industry": "",
        "status": "benchmark",
        "score": np.nan,
        "signal_date": signal_date.strftime("%Y-%m-%d"),
    }

    bench_row.update(bench_path)
    detail_rows.append(bench_row)

    skipped = 0

    for _, stock in selected_df.iterrows():
        code = str(stock["code"])

        if code not in stock_data:
            skipped += 1
            continue

        path = calc_forward_return_path_by_calendar(
            df=stock_data[code],
            forward_dates=forward_dates,
            buy_price_field=BUY_PRICE_FIELD,
        )

        if not path:
            skipped += 1
            continue

        row = {
            "row_type": "stock",
            "code": code,
            "name": stock.get("name", code_name_map.get(code, "")),
            "industry": stock.get("industry", ""),
            "status": stock["status"],
            "score": stock["score"],
            "score_trend_status": stock.get("score_trend_status", ""),
            "score_change_5d": stock.get("score_change_5d", np.nan),
            "trend_confirmed_by_score": stock.get("trend_confirmed_by_score", False),
            "signal_date": signal_date.strftime("%Y-%m-%d"),
            "close_on_signal": stock["close"],
            "ret20": stock["ret20"],
            "ret60": stock["ret60"],
            "ret20_rank_pct": stock["ret20_rank_pct"],
            "ret60_rank_pct": stock["ret60_rank_pct"],
            "amount_ma20": stock["amount_ma20"],
            "amount_ratio_5_20": stock["amount_ratio_5_20"],
            "max_dd20": stock["max_dd20"],
            "atr20_pct": stock["atr20_pct"],
            "dist_to_60d_high": stock["dist_to_60d_high"],
        }

        row.update(path)
        detail_rows.append(row)

    detail_df = pd.DataFrame(detail_rows)

    if len(detail_df) <= 1:
        print("入选股票未来数据不足，无法生成走势分析。")
        return

    if skipped > 0:
        print(f"有 {skipped} 只股票由于涨停买不进、买入日无交易、价格异常或未来数据缺失被跳过。")

    group_summary_df = build_forward_summary(
        detail_df=detail_df,
        forward_days=forward_days,
        group_col="status",
    )

    overall_summary_df = build_overall_forward_summary(
        detail_df=detail_df,
        forward_days=forward_days,
    )

    output_xlsx = os.path.join(
        OUTPUT_DIR,
        f"{FORWARD_ANALYSIS_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}.xlsx",
    )

    output_detail_csv = os.path.join(
        OUTPUT_DIR,
        f"{FORWARD_ANALYSIS_OUTPUT_PREFIX}_{signal_date.strftime('%Y-%m-%d')}_detail.csv",
    )

    try:
        save_forward_excel(
            output_xlsx=output_xlsx,
            signal_date=signal_date,
            detail_df=detail_df,
            group_summary_df=group_summary_df,
            overall_summary_df=overall_summary_df,
            selected_df=selected_df,
            pool_df=pool_df,
            forward_days=forward_days,
        )

        print(f"\n友好版分析结果已保存到：{output_xlsx}")

    except Exception as e:
        print(f"保存友好版 Excel 失败：{e}")
        detail_df.to_csv(output_detail_csv, index=False, encoding="utf-8-sig")
        print(f"走势明细已保存到：{output_detail_csv}")

    print("\n========== 未来走势摘要 ==========")

    key_horizons = [1, 3, 5, 10, 20, 30]
    key_horizons = [h for h in key_horizons if h <= forward_days]

    if not overall_summary_df.empty:
        display_df = overall_summary_df[
            overall_summary_df["horizon"].isin(key_horizons)
        ].copy()

        if not display_df.empty:
            print("\n固定持有整体统计：")
            print(display_df.to_string(index=False, formatters={
                "avg_ret": "{:.2%}".format,
                "median_ret": "{:.2%}".format,
                "win_rate": "{:.2%}".format,
                "bench_ret": "{:.2%}".format,
                "avg_excess": "{:.2%}".format,
                "outperform_rate": "{:.2%}".format,
            }))

    if not group_summary_df.empty:
        display_group_df = group_summary_df[
            group_summary_df["horizon"].isin(key_horizons)
        ].copy()

        if not display_group_df.empty:
            print("\n固定持有分组统计：")
            print(display_group_df.to_string(index=False, formatters={
                "avg_ret": "{:.2%}".format,
                "median_ret": "{:.2%}".format,
                "win_rate": "{:.2%}".format,
                "bench_ret": "{:.2%}".format,
                "avg_excess": "{:.2%}".format,
                "outperform_rate": "{:.2%}".format,
            }))


# =========================================================
# 11. 命令行入口
# =========================================================

def parse_args():
    parser = argparse.ArgumentParser(
        description="增强版：指定日期股票分类 / 评分趋势 / 短线强势池 / 未来走势分析"
    )

    parser.add_argument(
        "--date",
        required=True,
        help="信号日期，例如 2026-06-18。",
    )

    parser.add_argument(
        "--days",
        type=int,
        default=30,
        help="统计未来多少个交易日，默认 30。分类模式下无效。",
    )

    parser.add_argument(
        "--classify-only",
        action="store_true",
        help="只对指定日期股票进行分类，不做未来走势回测。",
    )

    parser.add_argument(
        "--statuses",
        type=str,
        default="短中线转强池,短线强势池,强趋势池,趋势观察池",
        help="要分析的股票池，逗号分隔。",
    )

    parser.add_argument(
        "--use-prev-trading-day",
        action="store_true",
        help="如果输入日期不是交易日，则自动使用前一个交易日。",
    )

    parser.add_argument(
        "--cache-dir",
        type=str,
        default="data",
        help="本地数据目录，默认 data。",
    )

    parser.add_argument(
        "--output-dir",
        type=str,
        default="output",
        help="输出目录，默认 output。",
    )

    parser.add_argument(
        "--universe-source",
        type=str,
        default="local",
        choices=["local", "ak", "auto"],
        help=(
            "股票池来源："
            "local=纯本地扫描 data/*.csv，不联网；"
            "ak=联网用 AkShare 获取代码名称并和本地 CSV 求交集；"
            "auto=优先 AkShare，失败后回退本地。默认 local。"
        ),
    )

    parser.add_argument(
        "--min-bars",
        type=int,
        default=130,
        help="最少K线数量，默认 130。",
    )

    parser.add_argument(
        "--min-score",
        type=int,
        default=75,
        help="趋势观察池最低评分，默认 75。",
    )

    parser.add_argument(
        "--strong-score",
        type=int,
        default=85,
        help="强趋势池最低评分，默认 85。",
    )

    parser.add_argument(
        "--min-amount",
        type=float,
        default=80_000_000,
        help="20日平均成交额门槛，默认 80000000。",
    )

    parser.add_argument(
        "--buy-price-field",
        type=str,
        default="open",
        choices=["open", "close"],
        help="买入基准价格字段，默认 open。",
    )

    parser.add_argument(
        "--score-trend-lookback",
        type=int,
        default=10,
        help="评分趋势回看交易日数量，默认 10。",
    )

    parser.add_argument(
        "--no-market-filter",
        action="store_true",
        help="关闭大盘环境过滤。默认开启。",
    )

    parser.add_argument(
        "--no-industry-filter",
        action="store_true",
        help="关闭行业强度过滤。默认开启，但需要 data/_stock_industry.csv。",
    )

    parser.add_argument(
        "--no-limit-check",
        action="store_true",
        help="关闭涨跌停不可交易模拟。默认开启。",
    )

    parser.add_argument(
        "--buy-commission",
        type=float,
        default=0.0003,
        help="买入佣金率，默认 0.0003。",
    )

    parser.add_argument(
        "--sell-commission",
        type=float,
        default=0.0003,
        help="卖出佣金率，默认 0.0003。",
    )

    parser.add_argument(
        "--stamp-tax",
        type=float,
        default=0.001,
        help="卖出印花税，默认 0.001。",
    )

    parser.add_argument(
        "--slippage",
        type=float,
        default=0.0005,
        help="滑点比例，默认 0.0005。",
    )

    return parser.parse_args()


def main():
    global CACHE_DIR
    global OUTPUT_DIR
    global BENCHMARK_CACHE_FILE
    global STOCK_NAME_FILE
    global STOCK_INDUSTRY_FILE

    global MIN_BARS
    global MIN_SCORE
    global STRONG_SCORE
    global MIN_AVG_AMOUNT_20
    global BUY_PRICE_FIELD

    global MARKET_FILTER_ENABLED
    global INDUSTRY_FILTER_ENABLED
    global LIMIT_CHECK_ENABLED

    global BUY_COMMISSION_RATE
    global SELL_COMMISSION_RATE
    global STAMP_TAX_RATE
    global SLIPPAGE_RATE

    global SCORE_TREND_LOOKBACK

    args = parse_args()

    CACHE_DIR = args.cache_dir
    OUTPUT_DIR = args.output_dir
    BENCHMARK_CACHE_FILE = os.path.join(CACHE_DIR, "_benchmark_000300.csv")
    STOCK_NAME_FILE = os.path.join(CACHE_DIR, "_stock_names.csv")
    STOCK_INDUSTRY_FILE = os.path.join(CACHE_DIR, "_stock_industry.csv")

    MIN_BARS = args.min_bars
    MIN_SCORE = args.min_score
    STRONG_SCORE = args.strong_score
    MIN_AVG_AMOUNT_20 = args.min_amount
    BUY_PRICE_FIELD = args.buy_price_field

    SCORE_TREND_LOOKBACK = args.score_trend_lookback

    MARKET_FILTER_ENABLED = not args.no_market_filter
    INDUSTRY_FILTER_ENABLED = not args.no_industry_filter
    LIMIT_CHECK_ENABLED = not args.no_limit_check

    BUY_COMMISSION_RATE = args.buy_commission
    SELL_COMMISSION_RATE = args.sell_commission
    STAMP_TAX_RATE = args.stamp_tax
    SLIPPAGE_RATE = args.slippage

    statuses = args.statuses.replace("，", ",")
    statuses = [s.strip() for s in statuses.split(",") if s.strip()]

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    print("========== 增强版：指定日期股票分类 / 评分趋势 / 未来走势分析 ==========")
    print(f"输入信号日期: {args.date}")
    print(f"运行模式: {'仅分类' if args.classify_only else '回测分析'}")
    print(f"未来交易日数: {args.days}")
    print(f"分析池: {statuses}")
    print(f"数据目录: {CACHE_DIR}")
    print(f"输出目录: {OUTPUT_DIR}")
    print(f"基准文件: {BENCHMARK_CACHE_FILE}")
    print(f"股票池来源: {args.universe_source}")
    print(f"最少K线数量: {MIN_BARS}")
    print(f"最低评分: {MIN_SCORE}")
    print(f"强趋势评分: {STRONG_SCORE}")
    print(f"评分趋势回看交易日: {SCORE_TREND_LOOKBACK}")
    print(f"20日平均成交额门槛: {MIN_AVG_AMOUNT_20:,.0f}")
    print(f"买入价格字段: {BUY_PRICE_FIELD}")

    print(f"大盘环境过滤: {'开启' if MARKET_FILTER_ENABLED else '关闭'}")
    print(f"行业强度过滤: {'开启' if INDUSTRY_FILTER_ENABLED else '关闭'}")
    print(f"涨跌停处理: {'开启' if LIMIT_CHECK_ENABLED else '关闭'}")

    print(f"买入佣金: {BUY_COMMISSION_RATE:.4%}")
    print(f"卖出佣金: {SELL_COMMISSION_RATE:.4%}")
    print(f"印花税: {STAMP_TAX_RATE:.4%}")
    print(f"滑点: {SLIPPAGE_RATE:.4%}")

    if args.universe_source == "local":
        print("注意：当前为纯本地模式，不联网。")
    elif args.universe_source == "ak":
        print("注意：当前会联网获取股票代码和名称，但不会下载历史行情。")
    elif args.universe_source == "auto":
        print("注意：当前会尝试联网获取股票代码和名称，失败后回退本地；不会下载历史行情。")

    print("====================================================================")

    print("\n构建股票池...")
    try:
        universe = get_universe(CACHE_DIR, source=args.universe_source)
    except Exception as e:
        print(f"构建股票池失败：{e}")
        return

    if universe is None or universe.empty:
        print("股票池为空，程序终止。")
        return

    codes = universe["code"].astype(str).apply(normalize_code).tolist()
    code_name_map = build_code_name_map(universe)
    industry_map = load_local_stock_industry(CACHE_DIR)

    print(f"最终可分析本地股票数量: {len(codes)}")
    print(f"有名称的股票数量: {sum(1 for c in codes if code_name_map.get(c, ''))}")
    print(f"有行业信息的股票数量: {sum(1 for c in codes if industry_map.get(c, ''))}")

    print("\n加载本地基准数据...")
    try:
        bench_df = fetch_benchmark_data()
    except Exception as e:
        print(f"加载基准数据失败：{e}")
        return

    try:
        signal_date = resolve_signal_date(
            requested_date=args.date,
            bench_df=bench_df,
            use_prev_trading_day=args.use_prev_trading_day,
        )
    except Exception as e:
        print(f"信号日期错误：{e}")
        return

    print("\n加载本地股票数据...")
    stock_data = preload_all_stock_data(codes)

    if not stock_data:
        print("没有成功加载任何股票数据，程序终止。")
        return

    if args.classify_only:
        analyze_pool_classification(
            signal_date=signal_date,
            stock_data=stock_data,
            bench_df=bench_df,
            code_name_map=code_name_map,
            industry_map=industry_map,
            score_trend_lookback=SCORE_TREND_LOOKBACK,
        )
    else:
        analyze_pool_forward_returns(
            signal_date=signal_date,
            stock_data=stock_data,
            bench_df=bench_df,
            code_name_map=code_name_map,
            industry_map=industry_map,
            forward_days=args.days,
            statuses=statuses,
            score_trend_lookback=SCORE_TREND_LOOKBACK,
        )


if __name__ == "__main__":
    main()